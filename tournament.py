#!/usr/bin/env python3
"""
Water Polo Tournament Translator
Shows your team's schedule, results, and bracket scenarios.

Usage:
    python3 tournament.py               → interactive team selector
    python3 tournament.py "TEAM NAME"   → jump straight to results
"""

import sys, re, os
import openpyxl
from datetime import datetime, date

XLSX_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "2026 TURBO OC CUP.xlsx")

SKIP_SHEETS = {
    "TEAM LISTING AND BRACKETS",
    "CHICLETS MASTER",
    "MASTER BY DIVISION",
    "MASTER BY LOCATION",
    "MASTER BY TIME",
}

# Regex to strip any bracket/pool prefix from a team slot
_PREFIX_RE = re.compile(
    r"^(?:"
    r"\d+(?:st|nd|rd|th)[A-Z]-"       # 1stA-, 2ndB-, etc.
    r"|[A-Z]\d+\([^)]+\)-"            # D1(1stA)-, E2(2ndB)-, etc.
    r"|[WL]#[^-]+-"                   # W#5-, L#B1/B4-, etc.
    r"|[A-Z]\d+-"                     # A1-, B3-, etc.
    r")(.*)",
    re.IGNORECASE,
)

# Placeholder slots that are NOT real team names
_PLACEHOLDER_RE = re.compile(
    r"^(?:"
    r"\d+(?:st|nd|rd|th)[A-Z]-?\s*$"  # "1stA-" alone
    r"|[WL]#"                          # W#... L#...
    r"|[A-Z]\d+\([^)]+\)-?\s*$"       # D1(1stA)- alone
    r")",
    re.IGNORECASE,
)

# ── Parsing ────────────────────────────────────────────────────────────────────

def load_workbook_data(filepath=XLSX_PATH):
    """Return (all_games, teams_by_division) from the spreadsheet."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_games = []
    teams_by_division = {}   # sheet_name → sorted list of team name strings

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        games = _parse_sheet(ws, sheet_name)
        all_games.extend(games)

        # Collect real team names for this division
        seen = set()
        for g in games:
            for slot in (g["white_team"], g["dark_team"]):
                name = strip_prefix(slot)
                if name and not _PLACEHOLDER_RE.match(name):
                    seen.add(name)
        if seen:
            teams_by_division[sheet_name] = sorted(seen)

    return all_games, teams_by_division


def _parse_sheet(ws, sheet_name):
    games = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        if len(row) < 8:
            continue
        date_val, time_val, location, game_id = row[0], row[1], row[2], row[3]
        white_team, white_score = row[4], row[5]
        dark_team,  dark_score  = row[6], row[7]
        comments = row[8] if len(row) > 8 else None
        division = row[9] if len(row) > 9 else None

        if not isinstance(date_val, datetime):
            continue
        if not game_id or not isinstance(game_id, str):
            continue
        if not white_team or not dark_team:
            continue

        games.append({
            "date":        date_val.date(),
            "time":        time_val if hasattr(time_val, "hour") else None,
            "location":    str(location).strip() if location else "TBD",
            "game_id":     game_id.strip(),
            "white_team":  str(white_team).strip(),
            "white_score": white_score,
            "dark_team":   str(dark_team).strip(),
            "dark_score":  dark_score,
            "comments":    str(comments).strip() if comments else "",
            "division":    str(division).strip() if division else sheet_name,
            "sheet":       sheet_name,
            "played":      white_score is not None and dark_score is not None,
        })
    return games


# ── Team name helpers ──────────────────────────────────────────────────────────

def strip_prefix(team_str: str) -> str:
    m = _PREFIX_RE.match(team_str.strip())
    return m.group(1).strip() if m else team_str.strip()


def team_matches(team_slot: str, search: str) -> bool:
    return search.upper() in strip_prefix(team_slot).upper()


def find_team_games(all_games, team_name, division_sheet=None):
    matches = [
        g for g in all_games
        if (division_sheet is None or g["sheet"] == division_sheet)
        and (team_matches(g["white_team"], team_name)
             or team_matches(g["dark_team"], team_name))
    ]
    matches.sort(key=lambda g: (g["date"] or date.min, g["time"] or datetime.min.time()))
    return matches


# ── Bracket tracing ────────────────────────────────────────────────────────────

def _game_number(game_id: str):
    m = re.search(r"(\d+)$", game_id)
    return str(int(m.group(1))) if m else None


def find_next_games(game, division_games):
    num = _game_number(game["game_id"])
    if num is None:
        return None, None
    winner_next = loser_next = None
    for g in division_games:
        if g["game_id"] == game["game_id"]:
            continue
        for slot in (g["white_team"], g["dark_team"]):
            m = re.match(r"^([WL])#(\w+)-", slot)
            if m:
                ref = re.search(r"(\d+)$", m.group(2))
                ref_num = str(int(ref.group(1))) if ref else m.group(2)
                if ref_num == num:
                    if m.group(1).upper() == "W":
                        winner_next = g
                    else:
                        loser_next = g
    return winner_next, loser_next


# ── Formatting helpers ─────────────────────────────────────────────────────────

def describe_slot(slot: str) -> str:
    """Human-readable opponent description from a raw team slot string."""
    m = re.match(
        r"^(?:"
        r"(\d+(?:st|nd|rd|th))([A-Z])-(.+)"       # 1stA-TEAM
        r"|([A-Z]\d+)\(([^)]+)\)-(.+)"             # D1(1stA)-TEAM
        r"|([WL])#([^-]+)-(.+)"                    # W#5-TEAM  or  L#B1/B4-TEAM
        r"|([A-Z])(\d+)-(.+)"                      # A1-TEAM
        r"|(.*)"                                    # plain name
        r")$",
        slot.strip(),
        re.IGNORECASE,
    )
    if not m:
        return slot.strip()

    # 1stA-TEAM
    if m.group(1):
        return f"{m.group(1)} place Pool {m.group(2)}  →  {m.group(3)}"
    # D1(1stA)-TEAM
    if m.group(4):
        return f"{m.group(5)} from {m.group(4)}  →  {m.group(6)}"
    # W#5-TEAM
    if m.group(7):
        wl = "Winner" if m.group(7).upper() == "W" else "Loser"
        return f"{wl} of game #{m.group(8)}  →  {m.group(9)}"
    # A1-TEAM
    if m.group(10):
        return m.group(12)
    # plain
    return m.group(13) or slot.strip()


def _fmt_time(t) -> str:
    return t.strftime("%I:%M %p").lstrip("0") if t else "TBD"

def _fmt_date(d) -> str:
    return d.strftime("%A, %b %-d") if d else "TBD"

def _fmt_score(game) -> str:
    ws, ds = game["white_score"], game["dark_score"]
    if ws is None or ds is None:
        return None
    try:
        return f"{int(float(ws))} – {int(float(ds))}"
    except (TypeError, ValueError):
        return f"{ws} – {ds}"

def _result(game, your_team) -> str:
    ws, ds = game["white_score"], game["dark_score"]
    if ws is None or ds is None:
        return ""
    try:
        ws, ds = float(ws), float(ds)
    except (TypeError, ValueError):
        return ""
    your_white = team_matches(game["white_team"], your_team)
    yours = ws if your_white else ds
    opp   = ds if your_white else ws
    if yours > opp:  return "✅ WIN"
    if yours < opp:  return "❌ LOSS"
    return "🤝 TIE"


def print_game_block(game, your_team, division_games=None):
    opp_slot  = game["dark_team"] if team_matches(game["white_team"], your_team) else game["white_team"]
    your_color = "WHITE" if team_matches(game["white_team"], your_team) else "DARK"
    opp_desc  = describe_slot(opp_slot)

    print(f"  📅  {_fmt_date(game['date'])}   {_fmt_time(game['time'])}")
    print(f"  📍  {game['location']}")
    print(f"  🆚  vs {opp_desc}   (you play {your_color})")

    score = _fmt_score(game)
    if score:
        print(f"  📊  {score}   {_result(game, your_team)}")
    else:
        print(f"  ⏳  Not yet played")

    if game["comments"]:
        print(f"  📝  {game['comments']}")

    # Bracket scenario
    if division_games and not game["played"]:
        winner_next, loser_next = find_next_games(game, division_games)
        if winner_next or loser_next:
            print()
            if winner_next:
                opp_w = describe_slot(
                    winner_next["dark_team"] if team_matches(winner_next["white_team"], your_team)
                    else winner_next["white_team"]
                )
                print(f"  ┌─ IF YOU WIN  →  vs {opp_w}")
                print(f"  │   {_fmt_date(winner_next['date'])}   {_fmt_time(winner_next['time'])}   {winner_next['location']}")
            else:
                print(f"  ┌─ IF YOU WIN  →  (no further game found)")

            if loser_next:
                opp_l = describe_slot(
                    loser_next["dark_team"] if team_matches(loser_next["white_team"], your_team)
                    else loser_next["white_team"]
                )
                print(f"  └─ IF YOU LOSE →  vs {opp_l}")
                print(f"      {_fmt_date(loser_next['date'])}   {_fmt_time(loser_next['time'])}   {loser_next['location']}")
            else:
                print(f"  └─ IF YOU LOSE →  (eliminated / no further game)")

    print()


# ── Team Selector ──────────────────────────────────────────────────────────────

def _clear():
    os.system("cls" if os.name == "nt" else "clear")


def _division_label(sheet_name: str) -> str:
    """Friendly label for a sheet name."""
    return sheet_name.replace("-", " ").replace("_", " ")


def select_team_interactive(teams_by_division):
    """
    Two-step interactive selector:
      1. Pick a division
      2. Pick a team within that division
    Returns (team_name, sheet_name) or None if user quits.
    """
    divisions = list(teams_by_division.keys())

    while True:
        _clear()
        print("╔══════════════════════════════════════════════════════════╗")
        print("║          🏊  TURBO OC CUP 2026  —  Team Finder          ║")
        print("╚══════════════════════════════════════════════════════════╝")
        print()
        print("  Select your division:\n")

        for i, div in enumerate(divisions, 1):
            team_count = len(teams_by_division[div])
            print(f"  {i:>2}.  {_division_label(div)}  ({team_count} teams)")

        print()
        print("  Q.  Quit")
        print()
        raw = input("  Enter number: ").strip().lower()

        if raw == "q":
            return None

        try:
            idx = int(raw) - 1
            if not (0 <= idx < len(divisions)):
                raise ValueError
        except ValueError:
            input("  ⚠️  Invalid choice. Press Enter to try again.")
            continue

        chosen_division = divisions[idx]
        teams = teams_by_division[chosen_division]

        # ── Step 2: pick team ──
        while True:
            _clear()
            print("╔══════════════════════════════════════════════════════════╗")
            print("║          🏊  TURBO OC CUP 2026  —  Team Finder          ║")
            print("╚══════════════════════════════════════════════════════════╝")
            print()
            print(f"  Division: {_division_label(chosen_division)}\n")
            print("  Select your team:\n")

            for i, team in enumerate(teams, 1):
                print(f"  {i:>2}.  {team}")

            print()
            print("  B.  ← Back to divisions")
            print("  Q.  Quit")
            print()
            raw2 = input("  Enter number: ").strip().lower()

            if raw2 == "q":
                return None
            if raw2 == "b":
                break

            try:
                tidx = int(raw2) - 1
                if not (0 <= tidx < len(teams)):
                    raise ValueError
            except ValueError:
                input("  ⚠️  Invalid choice. Press Enter to try again.")
                continue

            return teams[tidx], chosen_division


# ── Output ─────────────────────────────────────────────────────────────────────

def show_team_results(all_games, team_name, sheet_name=None):
    team_games = find_team_games(all_games, team_name, division_sheet=sheet_name)

    if not team_games:
        print(f"\n❌  No games found for '{team_name}'.")
        return

    divisions = sorted({g["division"] for g in team_games})
    played   = [g for g in team_games if g["played"]]
    upcoming = [g for g in team_games if not g["played"]]

    division_games_map = {}
    for g in all_games:
        division_games_map.setdefault(g["sheet"], []).append(g)

    bar = "═" * 62
    print()
    print(bar)
    print(f"  🏊  {team_name.upper()}")
    print(f"  Division: {', '.join(divisions)}")
    print(f"  {len(played)} result(s)   •   {len(upcoming)} upcoming game(s)")
    print(bar)
    print()

    if played:
        print("▶  RESULTS")
        print("─" * 62)
        for game in played:
            print_game_block(game, team_name)

    if upcoming:
        print("▶  UPCOMING GAMES")
        print("─" * 62)
        for game in upcoming:
            dg = division_games_map.get(game["sheet"], [])
            print_game_block(game, team_name, division_games=dg)

    if not played and not upcoming:
        print("  (No game data found — check team name.)\n")


# ── Entry point ────────────────────────────────────────────────────────────────

def main():
    print("Loading tournament data…", end=" ", flush=True)
    all_games, teams_by_division = load_workbook_data()
    print(f"{len(all_games)} games loaded.\n")

    # Command-line shortcut: python3 tournament.py "TEAM NAME"
    if len(sys.argv) > 1:
        team_name = " ".join(sys.argv[1:])
        show_team_results(all_games, team_name)
        return

    # Interactive selector loop — lets multiple people use it back-to-back
    while True:
        result = select_team_interactive(teams_by_division)
        if result is None:
            print("\nGoodbye! 🏊\n")
            break

        team_name, sheet_name = result
        _clear()
        show_team_results(all_games, team_name, sheet_name=sheet_name)

        print("\n" + "─" * 62)
        again = input("  Press Enter to look up another team, or Q to quit: ").strip().lower()
        if again == "q":
            print("\nGoodbye! 🏊\n")
            break


if __name__ == "__main__":
    main()
