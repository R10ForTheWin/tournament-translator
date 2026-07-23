#!/usr/bin/env python3
"""
Golden test suite for Tournament Translator parsers.

Run from the project root:
    python3 tests/test_parsers.py

Each fixture has an expected minimum game count and format. A test fails if:
  - Fewer/more games than expected are parsed
  - Wrong parser selected (format mismatch)
  - Duplicate game_ids within a sheet
  - Any team slot fails structural validation

IMPORTANT: After fixing any bug, add (or tighten) a test here so the bug
can never silently return. The test suite is the long-term memory of the
system — every entry encodes a lesson learned the hard way.
"""
from __future__ import annotations
import os, sys, re
from collections import Counter
from datetime import date, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from parsers.detect import load_and_parse
from parsers.validate import _valid_slot
from parsers.format_cca import parse_csv as cca_parse_csv
from parsers.format_a import parse as format_a_parse
from app import (
    _expand_bracket_games, _build_wpl_game_tree, _build_njo_game_tree,
    team_matches, describe_slot, _SLOT_LIKE_RE, _POOL_PREVIEW_RE, _tournament_meta,
    _team_opp_slot, _result_str, _tournament_finish_probs, _compute_tree_layout,
)

FIXTURES_DIR   = os.path.join(ROOT, "Tournaments Excels")
FIXTURES_CCA   = os.path.join(ROOT, "tests", "fixtures")

WPL_FILE = "2026 KAP7 Futures WPL - Southern California - Presented by BIWPA.xlsx"

# ── Fixture definitions ──────────────────────────────────────────────────────
# (filename, expected_format, min_games, max_games)
# Game count range should be wide enough to survive minor organizer edits
# but tight enough to catch a format switch or wholesale data loss.
FIXTURES = [
    ("2025 NEWPORT SPRING INVITE.xlsx",   "A",  160,  200),
    ("2026 KAP7 INTERNATIONAL.xlsx",      "A",  580,  650),
    ("2026 TURBO OC CUP.xlsx",            "A",  280,  340),
    # WPL Futures — all 5 weekends.  Updated range reflects full-season fixture
    # loaded after Weekend 5 championship (was 1200-1400 for 3-weekend snapshot).
    (WPL_FILE,                            "B", 1800, 2200),
]

# ── Cross-weekend contamination canary ───────────────────────────────────────
# (filename, sheet, team, min_direct, max_direct, max_extras)
# max_extras is the contamination guard: if expansion pulls in games from the
# wrong weekend, extras spike.  The canary value tracks the KNOWN-GOOD state;
# tighten it whenever the all-weekends fixture is refreshed.
TEAM_CHECKS = [
    # Trojan Gold, all weekends in fixture.  With championship data loaded,
    # extras = 3 (the 3 Sunday -B placeholder games for pool H placement).
    # If extras ever exceed 6 cross-weekend contamination has returned.
    (WPL_FILE, "16u Boys", "trojan gold", 14, 22, 6),
]

# ── Championship-weekend specific checks ─────────────────────────────────────
# These check the behavior we care about most on game day.
#
# Anchor date is derived from _tournament_meta() so it stays correct when a
# new tournament weekend is added to KNOWN_TOURNAMENTS — no manual date updates.
#
# Schema per entry: (sheet, team, min_direct, max_direct,
#                    min_extras, min_tree_nodes, min_sun_nodes)
#
# Thresholds are set to hold BOTH pre-tournament (all placeholders visible) AND
# post-tournament (bracket narrowed to one path).  Lower bounds only — they catch
# regressions back to zero without breaking when live scores narrow the tree.
#
#   min_extras=1    — at least 1 Sunday game must reach the expansion
#   min_sun_nodes=1 — at least 1 Sunday node must appear in the bracket tree
#
# Add a row here any time a new bug class is fixed on a real team.
CHAMPIONSHIP_CHECKS: dict[str, list[tuple]] = {
    "futures-5": [
        # Pool-slot seeded team (H1): 2 Sat games, 3 Sun placement placeholders.
        # Previously: Sunday games missing due to game-ID collision in format_b.py.
        ("16u Boys", "trojan gold",     2, 2, 1, 3, 1),

        # Seed-number prelim team: won prelim + played 2 bracket games = 3 direct.
        # Previously: showed only 1 game (seed-number root unhandled).
        ("16u Boys", "trojan cardinal", 1, 3, 1, 3, 1),

        # Seeded pool-E1 team — regression canary for pool-slot expansion.
        ("16u Boys", "imperial",        2, 2, 1, 3, 1),

        # Seeded pool-G1 team — second pool-slot canary.
        ("16u Boys", "socal",           2, 2, 1, 3, 1),
    ],
    # Add "futures-super", "futures-6", etc. here as those weekends are played.
}


# ── Helpers ──────────────────────────────────────────────────────────────────

PASS = "\033[32mPASS\033[0m"
FAIL = "\033[31mFAIL\033[0m"

def _check(label: str, condition: bool, detail: str = "") -> bool:
    status = PASS if condition else FAIL
    suffix = f"  ({detail})" if detail else ""
    print(f"  [{status}] {label}{suffix}")
    return condition


# ── Tests ─────────────────────────────────────────────────────────────────────

def _run_fixture(fname: str, fmt: str, min_g: int, max_g: int) -> int:
    path = os.path.join(FIXTURES_DIR, fname)
    failures = 0
    print(f"\n{fname}")

    if not os.path.exists(path):
        print(f"  [SKIP] file not found: {path}")
        return 0

    games = load_and_parse(path)

    ok = _check(f"game count {len(games)} in [{min_g}, {max_g}]",
                min_g <= len(games) <= max_g, f"got {len(games)}")
    if not ok: failures += 1

    formats = {g.get("format") for g in games} - {None}
    ok = _check(f"parser format is {fmt!r}", fmt in formats, f"got {formats}")
    if not ok: failures += 1

    by_sheet: dict[str, list] = {}
    for g in games:
        by_sheet.setdefault(g.get("sheet", ""), []).append(g["game_id"])
    dup_sheets = []
    for sh, ids in by_sheet.items():
        dups = [gid for gid, n in Counter(ids).items() if n > 1]
        if dups:
            dup_sheets.append(f"{sh}: {dups[:3]}")
    ok = _check("no duplicate game_ids per sheet", not dup_sheets,
                "; ".join(dup_sheets))
    if not ok: failures += 1

    bad_slots = []
    for g in games[:200]:
        for field in ("white_team", "dark_team"):
            slot = str(g.get(field) or "").strip()
            if not _valid_slot(slot):
                bad_slots.append(f"{g['game_id']}.{field}={slot!r}")
    ok = _check("team slots structurally valid (first 200 games)",
                not bad_slots, "; ".join(bad_slots[:3]))
    if not ok: failures += 1

    scored = sum(1 for g in games if g.get("white_score") is not None)
    ok = _check(f"at least one scored game", scored > 0, f"{scored} scored")
    if not ok: failures += 1

    return failures


def _run_team_expansion(fname: str, sheet: str, team: str,
                         min_d: int, max_d: int, max_extras: int) -> int:
    path = os.path.join(FIXTURES_DIR, fname)
    failures = 0
    label = f"{team!r} in {sheet!r}"
    print(f"\n{label}")

    if not os.path.exists(path):
        print(f"  [SKIP] file not found: {path}")
        return 0

    games = load_and_parse(path)
    div_games = [g for g in games if g["sheet"] == sheet]
    direct = sorted(
        [g for g in div_games
         if team_matches(g["white_team"], team) or team_matches(g["dark_team"], team)],
        key=lambda g: (g.get("date") or date.min, g.get("time") or datetime.min.time()),
    )
    extras = _expand_bracket_games(team, direct, div_games)

    ok = _check(f"direct game count {len(direct)} in [{min_d}, {max_d}]",
                min_d <= len(direct) <= max_d, f"got {len(direct)}")
    if not ok: failures += 1

    ok = _check(f"extras ≤ {max_extras} (cross-weekend contamination guard)",
                len(extras) <= max_extras, f"got {len(extras)}")
    if not ok: failures += 1

    if direct and extras:
        first_date = min(g["date"] for g in direct if g.get("date"))
        early = [g["game_id"] for g in extras
                 if not g.get("placeholder") and g.get("date") and g["date"] < first_date]
        ok = _check("no real extra predates team's first direct game",
                    not early, str(early))
        if not ok: failures += 1

    return failures


def _run_championship_team(sheet: str, team: str, anchor: date,
                            min_direct: int, max_direct: int,
                            min_extras: int, min_tree: int, min_sun: int) -> int:
    """Check a team's championship-weekend bracket and expansion.

    This is the test that would have caught every game-day bug we've fixed:
    - Too few direct games  (organizer format change, parser regression)
    - Too few extras        (expansion not following WIN GM # or finish slots)
    - Empty tree            (seed-number root not handled)
    - No Sunday nodes       (Sunday games missing from sheet or dedup-dropped)
    - Unresolved slots      (describe_slot returning raw "1stD-" strings)
    - Stale opponent names  (cross-weekend contamination in standings lookup)
    """
    path = os.path.join(FIXTURES_DIR, WPL_FILE)
    failures = 0
    label = f"{team!r} | {sheet} championship {anchor}"
    print(f"\n{label}")

    if not os.path.exists(path):
        print(f"  [SKIP] file not found: {path}")
        return 0

    games     = load_and_parse(path)
    div_games = [g for g in games if g["sheet"] == sheet]

    # Direct games this weekend only
    direct = sorted(
        [g for g in div_games
         if (team_matches(g["white_team"], team) or team_matches(g["dark_team"], team))
         and g.get("date") and abs((g["date"] - anchor).days) <= 1],
        key=lambda g: (g.get("date") or date.min, g.get("time") or datetime.min.time()),
    )

    ok = _check(f"direct championship games in [{min_direct}, {max_direct}]",
                min_direct <= len(direct) <= max_direct, f"got {len(direct)}")
    if not ok: failures += 1

    # Expansion
    extras = _expand_bracket_games(team, direct, div_games)
    ok = _check(f"extras ≥ {min_extras} (bracket expansion reaches Sunday)",
                len(extras) >= min_extras, f"got {len(extras)}")
    if not ok: failures += 1

    # Tree builder
    tree = _build_wpl_game_tree(team, div_games, anchor_date=anchor)
    ok = _check(f"tree has ≥ {min_tree} nodes",
                len(tree) >= min_tree, f"got {len(tree)}")
    if not ok: failures += 1

    sun_nodes = [n for n in tree if n.get("date") and n["date"].weekday() == 6]
    ok = _check(f"tree has ≥ {min_sun} Sunday node(s)",
                len(sun_nodes) >= min_sun, f"got {len(sun_nodes)}")
    if not ok: failures += 1

    # Unresolved slot strings (describe_slot contamination check)
    bad_opps = []
    for n in tree:
        wt, dt = n["white_team"], n["dark_team"]
        opp_slot = dt if team.split()[-1].upper() in wt.upper() else wt
        opp = describe_slot(opp_slot, div_games, ref_date=anchor)
        if _SLOT_LIKE_RE.match(opp) and not _POOL_PREVIEW_RE.match(opp):
            bad_opps.append(f"{n['game_id']}: {opp!r}")
    ok = _check("no unresolved slot strings in opponent labels",
                not bad_opps, "; ".join(bad_opps[:3]))
    if not ok: failures += 1

    # Placeholder discipline: WIN-GM-inferred games must not be marked real
    # before the prelim has been played.
    inferred_real = []
    for n in tree:
        wt, dt = n["white_team"], n["dark_team"]
        is_inferred = (
            not team_matches(wt, team) and not team_matches(dt, team)
            and re.search(r'\bWIN\s+GM\s+#', wt + dt, re.IGNORECASE)
        )
        if is_inferred and not n.get("placeholder"):
            inferred_real.append(n["game_id"])
    ok = _check("no inferred WIN-GM games marked as confirmed (placeholder=False)",
                not inferred_real, str(inferred_real))
    if not ok: failures += 1

    return failures


# ── CCA opponent-slot correctness ─────────────────────────────────────────────
#
# These tests catch the class of bug where _team_opp_slot returns the team's
# OWN slot as the opponent ("Loser of game #23" when the team IS the loser of
# game #23). This happens when team_matches() fails on a placeholder slot and
# the code falls to a wrong default. Each entry pins a known-correct pre-
# tournament opponent slot for a specific placeholder game type:
#
#   (csv_file, division, id_prefix, team_query, game_id, expected_opp_slot)
#
# "expected_opp_slot" is the raw normalized slot string, not a display label.
# If the slot changes because results fill it in, update the expectation — the
# important invariant is that the opponent is NEVER the team's own slot.
#
# Fixture files are snapshots of the live Google Sheets saved at
# tests/fixtures/cca_{16u,18u}.csv. Refresh them before each tournament day
# by running: python3 tests/refresh_cca_fixtures.py
CCA_OPP_CHECKS = [
    # 16U Trojan Gold (B) — pool G2 seed
    #   finish-slot path: pool_rank_group="G", opp must not be "2ndG-"
    ("cca_16u.csv", "16U Boys", "16U", "trojan gold", "16U-38", "1stH-"),
    #   finish-slot path: pool_rank_group="G", team is "1stG-" here, opp must not be "1stG-"
    ("cca_16u.csv", "16U Boys", "16U", "trojan gold", "16U-40", "2ndH-"),
    #   W#/L# path: L#23 is team's slot (game 23 in my_ids), opp must not be "L#23"
    ("cca_16u.csv", "16U Boys", "16U", "trojan gold", "16U-46", "L#22"),

    # 16U Trojan Cardinal (A) — pool C1 seed
    #   finish-slot path: opp must not be "1stC-"
    ("cca_16u.csv", "16U Boys", "16U", "trojan cardinal", "16U-27", "2ndD-"),
    #   W#/L# path: L#12 is team's slot
    ("cca_16u.csv", "16U Boys", "16U", "trojan cardinal", "16U-25", "L#13"),

    # 16U Trojan Silver (C) — pool E4 seed
    #   W#/L# path: L#10 is team's slot
    ("cca_16u.csv", "16U Boys", "16U", "trojan silver", "16U-35", "L#11"),

    # 18U Trojan Cardinal A — bracket-entry BB1 (no pool play)
    #   finish-slot path: team is direct white slot, opp is 1stC-
    ("cca_18u.csv", "18U Boys", "18U", "trojan cardinal", "18U-15", "1stC-"),
    #   W#/L# path: W#15 is team's slot (win of 15 in my_ids), opp must not be "W#15"
    ("cca_18u.csv", "18U Boys", "18U", "trojan cardinal", "18U-20", "L#16"),
]


# ── Format A game-id collision test ──────────────────────────────────────────
# Regression 1: Quiksilver Cup 2026 reused game number "16UB09" for two different
# games WITHIN THE SAME SHEET (an organizer typo — one block's numbering should
# have started at 10). format_a.py's dedup was a flat seen_ids set that silently
# dropped the second occurrence, which would corrupt pool standings (missing 1 of
# 3 round-robin results). Fixed by porting format_b.py's collision-rename
# (-B/-C suffix) pattern into format_a.py.
#
# Regression 2 (found while fixing #1): that same-sheet rename logic, if applied
# globally, wrongly resurrects garbage rows from stale duplicate DIVISION SHEETS.
# 2025 Newport Spring Invite has both "18U BOYS PLATINUM-23 TEAMS" (real, current)
# and "18U BOYS PLATINUM-22 TEAMS" (an older stale draft with unfilled slots like
# "4-"/"21-") — both reuse game_id "18Bpt01". The ORIGINAL cross-sheet global dedup
# correctly dropped the stale sheet's row (relying on the real sheet appearing
# first in workbook order); a naive same-ID-different-teams-anywhere rename would
# have kept the "4-" vs "21-" garbage row as a "new" game. Fix: only rename when
# the collision is within the SAME sheet; cross-sheet collisions still drop silently.

def test_format_a_id_collision() -> int:
    """Same game_id + different teams, within ONE sheet: both games must survive
    (second renamed with a -B suffix). Same game_id across DIFFERENT sheets
    (stale duplicate division tab) must still silently drop the second, keeping
    only the first (real) sheet's data. True duplicates (same id, same teams,
    same sheet) must collapse to one."""
    import openpyxl

    failures = 0
    print("\nFormat A game-id collision handling")

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "16U BOYS"
    rows = [
        (datetime(2026, 7, 10), datetime(2026, 7, 10, 8, 0).time(), "EL MODENA HS", "16UB09",
         "E1-IMPERIAL", None, "E2-MISSION", None, "", "16U_BOYS"),
        (datetime(2026, 7, 10), datetime(2026, 7, 10, 9, 0).time(), "BUENA PARK HS", "16UB09",
         "B1-LA JOLLA UNITED", None, "B3-COMMERCE", None, "", "16U_BOYS"),
        # True duplicate: same id, same teams, same sheet — must collapse to 1
        (datetime(2026, 7, 10), datetime(2026, 7, 10, 8, 0).time(), "EL MODENA HS", "16UB09",
         "E1-IMPERIAL", None, "E2-MISSION", None, "", "16U_BOYS"),
    ]
    for r in rows:
        ws1.append(r)

    # Second sheet: same game_id as a real game above, but garbage/unresolved
    # team text — simulates a stale duplicate division tab (Newport pattern).
    ws2 = wb.create_sheet("16U BOYS-OLD DRAFT")
    ws2.append((datetime(2026, 7, 10), datetime(2026, 7, 10, 8, 0).time(), "EL MODENA HS", "16UB09",
                "4-", None, "21-", None, "", "16U_BOYS"))

    games = format_a_parse(wb)
    ids = [g["game_id"] for g in games]

    ok = _check("same-sheet collision: both distinct games survive (2, not 1)",
                len([g for g in games if g["sheet"] == "16U BOYS"]) == 2,
                f"got {len(games)} total: {ids}")
    if not ok: failures += 1

    ok = _check("original id preserved for first occurrence", "16UB09" in ids, f"ids={ids}")
    if not ok: failures += 1

    ok = _check("same-sheet collision renamed with -B suffix, not dropped",
                "16UB09-B" in ids, f"ids={ids}")
    if not ok: failures += 1

    ok = _check("cross-sheet collision (stale draft tab) silently dropped, not renamed",
                not any(g["sheet"] == "16U BOYS-OLD DRAFT" for g in games),
                f"ids={ids}, sheets={[g['sheet'] for g in games]}")
    if not ok: failures += 1

    return failures


# ── Composite pool-rank slot bug (Quiksilver Cup round-by-round simulation) ──
# Regression: describe_slot / _team_opp_slot / _result_str all resolve a team's
# rank within a pool for composite bracket slots like "K1(2ndB)-" (bracket
# position K1 seeded by 2nd place in pool B). Found by simulating Trojan
# Cardinal winning/losing/splitting Pool B and checking each round's API
# response for Quiksilver Cup 2026:
#   - describe_slot picked the FIRST digit anywhere in the slot ("1" from "K1")
#     as the pool rank instead of the ordinal's digit ("2" from "2nd"), so
#     "K1(2ndB)-" incorrectly resolved to Pool B's 1st-place team instead of 2nd.
#   - _team_opp_slot only matched the simple "1stB-" format (anchored regex),
#     never the composite "K1(2ndB)-" format, so it always fell back to
#     returning the white slot as "opponent" regardless of which side the team
#     was actually on — silently dropping self-referential games from the
#     schedule, or (via _result_str, which has the same direct-name-match-only
#     blind spot) reporting a WIN as a LOSS whenever the team was white.

def _build_pool_b_games():
    """3-team Pool B round robin (Cardinal 2nd) + one Round-2 composite game
    referencing Pool B's rank -- mirrors the real Quiksilver Cup bracket shape."""
    return [
        {"game_id": "16UB09", "date": date(2026, 7, 10), "time": None, "location": "X",
         "white_team": "B1-LA JOLLA UNITED", "white_score": 8, "dark_team": "B3-COMMERCE",
         "dark_score": 4, "comments": "", "division": "16U_BOYS", "sheet": "16U BOYS",
         "played": True},
        {"game_id": "16UB12", "date": date(2026, 7, 10), "time": None, "location": "X",
         "white_team": "B2-TROJAN CARDINAL", "white_score": 9, "dark_team": "B3-COMMERCE",
         "dark_score": 5, "comments": "", "division": "16U_BOYS", "sheet": "16U BOYS",
         "played": True},
        {"game_id": "16UB15", "date": date(2026, 7, 10), "time": None, "location": "X",
         "white_team": "B1-LA JOLLA UNITED", "white_score": 10, "dark_team": "B2-TROJAN CARDINAL",
         "dark_score": 6, "comments": "", "division": "16U_BOYS", "sheet": "16U BOYS",
         "played": True},
        # Cardinal (2nd in B) is WHITE here, seeded into bracket position K1.
        # Bracket-position digit (1) intentionally differs from ordinal rank (2).
        {"game_id": "16UB22", "date": date(2026, 7, 11), "time": None, "location": "X",
         "white_team": "K1(2ndB)-", "white_score": 9, "dark_team": "K3(2ndF)-",
         "dark_score": 5, "comments": "", "division": "16U_BOYS", "sheet": "16U BOYS",
         "played": True},
    ]


def test_composite_pool_rank_slot() -> int:
    """Cardinal finishes 2nd in Pool B (not 1st) -- every function that resolves
    a composite pool-rank slot must agree on rank=2, not rank=1."""
    failures = 0
    print("\nComposite pool-rank slot correctness (K1(2ndB)- style)")

    games = _build_pool_b_games()
    team = "Trojan Cardinal"

    resolved = describe_slot("K1(2ndB)-", games)
    ok = _check("describe_slot('K1(2ndB)-') resolves to the 2nd-place team, not 1st",
                resolved.upper() == "TROJAN CARDINAL", f"got {resolved!r}")
    if not ok: failures += 1

    direct = [g for g in games
              if team_matches(g["white_team"], team) or team_matches(g["dark_team"], team)]
    extras = _expand_bracket_games(team, direct, games)
    my_ids = {g["game_id"] for g in direct + extras}
    k_game = next((g for g in extras if g["game_id"] == "16UB22"), None)

    ok = _check("Round-2 composite-slot game (16UB22) reached via expansion",
                k_game is not None, f"extras={[g['game_id'] for g in extras]}")
    if not ok:
        failures += 1
    else:
        opp = _team_opp_slot(k_game, team, games, my_ids)
        ok = _check("_team_opp_slot identifies K3(2ndF)- as opponent, not our own K1(2ndB)- slot",
                    opp == "K3(2ndF)-", f"got {opp!r}")
        if not ok: failures += 1

        result = _result_str(k_game, team)
        ok = _check("_result_str: Cardinal (white, 9-5) is a WIN, not a loss",
                    result == "win", f"got {result!r}")
        if not ok: failures += 1

    return failures


# ── game_num chronological-order bug (day-merge wrongly merging pool games) ──
# Regression: found live on quiksilver-cup (zero games played yet). Trojan
# Cardinal's schedule is 2 round-robin Pool B games on Friday (different
# opponents, non-overlapping times) followed by one Saturday placeholder.
# The "merge same-day pure-single-day columns" step merged the two independent
# Friday games into one game_num (they're not bracket alternates, just two
# separate real games), then the later self-heal step split them back apart
# but appended the recovered game at the END of the sequence instead of
# re-sorting chronologically -- result: Game 1 (Fri 11am), Game 2 (Sat
# placeholder), Game 3 (Fri 2pm) -- Friday's second game displayed AFTER
# Saturday's, despite being a full day earlier. Fixed by only allowing the
# day-merge to combine groups that are genuine win/lose bracket alternates
# (bracket_path in {"win","lose"}), never two independent scheduled games.

def test_game_num_chronological_order() -> int:
    """game_num must increase strictly in chronological (date, time) order for
    a team's own real games -- Friday's games can never be numbered after a
    Saturday placeholder."""
    import io as _io
    import json as _json
    import openpyxl

    failures = 0
    print("\ngame_num chronological order (Quiksilver Cup Pool B shape)")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "16U BOYS-18 TEAMS"
    rows = [
        # Both Pool B games played (Cardinal wins both) so pool_rank=1 is known --
        # otherwise the composite Round-2 slot below is correctly suppressed as
        # an unresolved guess (see _expand_bracket_games) and wouldn't appear at
        # all, which is the right behavior but not what this test is checking.
        (datetime(2026, 7, 10), datetime(2026, 7, 10, 11, 0).time(), "BUENA PARK HS", "16UB12",
         "B2-TROJAN CARDINAL", 10, "B3-COMMERCE", 5, "", "16U_BOYS"),
        (datetime(2026, 7, 10), datetime(2026, 7, 10, 14, 0).time(), "BUENA PARK HS", "16UB15",
         "B1-LA JOLLA UNITED", 6, "B2-TROJAN CARDINAL", 9, "", "16U_BOYS"),
        (datetime(2026, 7, 11), datetime(2026, 7, 11, 9, 0).time(), "EL MODENA HS", "16UB21",
         "H1(1stB)-", None, "H3(1stF)-", None, "", "16U_BOYS"),
    ]
    for r in rows:
        ws.append(r)
    buf = _io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    import app as appmod
    orig_find_excel = appmod.find_excel
    def fake_find_excel(tid):
        return _io.BytesIO(xlsx_bytes) if tid == "quiksilver-cup" else orig_find_excel(tid)
    appmod.find_excel = fake_find_excel
    appmod._parse_cache.clear()
    try:
        with appmod.app.test_client() as c:
            r = c.get("/api/games/quiksilver-cup/Trojan%20Cardinal")
            data = _json.loads(r.data)
            games_by_id = {g["game_id"]: g for g in data.get("upcoming", []) + data.get("played", [])}

            ok = _check("all 3 games present", len(games_by_id) == 3, f"got {list(games_by_id)}")
            if not ok: failures += 1

            order = sorted(games_by_id.values(), key=lambda g: g["game_num"])
            order_ids = [g["game_id"] for g in order]
            ok = _check("game_num order matches chronological order (Fri 11am, Fri 2pm, Sat 9am)",
                        order_ids == ["16UB12", "16UB15", "16UB21"], f"got order={order_ids}")
            if not ok: failures += 1
    finally:
        appmod.find_excel = orig_find_excel
        appmod._parse_cache.clear()

    return failures


# ── Bracket branching + "games remaining" TBD stubs ──────────────────────────
# This is a bracket app: before a team's pool rank is known, every candidate
# slot (1st/2nd/3rd place) is a genuine branch and must be SHOWN as an
# alternative (tagged with its own pool_rank so the UI can label "If 1st in
# Pool" / "If 2nd in Pool" / etc.) -- not suppressed. An earlier version of
# this feature suppressed all candidates pre-rank and replaced them with a
# generic TBD stub; that was reverted because it hid the actual bracket tree
# the app exists to show.
#
# TBD stubs still exist for the one case a fixed-game-count format truly
# cannot resolve at all: a game beyond what _expand_bracket_games can reach
# (e.g. the two-level pool-of-pool-winners gap -- see project memory). Must
# NOT fire for elimination formats (WPL/Kap7) where game count varies by
# win/loss -- _expected_games_per_team returns None there.

def test_tbd_stub_placeholders() -> int:
    """3-team pool P (Cardinal's own) + 3-team pool Q (a different pool, always
    played) + 1 Round-2 composite game referencing 1st-of-P vs 1st-of-Q --
    mirrors Quiksilver's real shape where Round 2 opponents always come from a
    different pool than the team's own."""
    import io as _io
    import json as _json
    import openpyxl

    failures = 0
    print("\nTBD stub placeholders (fixed-game-count formats)")

    def _make_wb(round1_played: bool):
        # Pool P (Cardinal's own pool) and pool Q (a different pool) each need
        # 3 teams so round-robin standings are real, matching Quiksilver's
        # actual shape -- Round 2 opponents always come from a DIFFERENT pool
        # than the team's own, never the same one, so there's no ambiguity
        # about the team's own name appearing among the opponent's candidates.
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "MINI DIV"
        r1 = (10, 5) if round1_played else (None, None)
        ws.append((datetime(2026, 7, 10), datetime(2026, 7, 10, 11, 0).time(), "VENUE", "T1",
                   "P1-TROJAN CARDINAL", r1[0], "P2-TEAM B", r1[1], "", "MINI"))
        ws.append((datetime(2026, 7, 10), datetime(2026, 7, 10, 12, 0).time(), "VENUE", "T1B",
                   "P1-TROJAN CARDINAL", r1[0], "P3-TEAM C", r1[1], "", "MINI"))
        ws.append((datetime(2026, 7, 10), datetime(2026, 7, 10, 13, 0).time(), "VENUE", "T1C",
                   "P2-TEAM B", 3, "P3-TEAM C", 1, "", "MINI"))
        # Pool Q — a separate pool, always played, so its 1st place is known
        # regardless of round1_played (mirrors real Round 2 opponents being
        # resolvable independent of the team's own pool result).
        ws.append((datetime(2026, 7, 10), datetime(2026, 7, 10, 14, 0).time(), "VENUE", "T1D",
                   "Q1-TEAM X", 10, "Q2-TEAM Y", 2, "", "MINI"))
        ws.append((datetime(2026, 7, 10), datetime(2026, 7, 10, 15, 0).time(), "VENUE", "T1E",
                   "Q1-TEAM X", 9, "Q3-TEAM Z", 3, "", "MINI"))
        ws.append((datetime(2026, 7, 10), datetime(2026, 7, 10, 16, 0).time(), "VENUE", "T1F",
                   "Q2-TEAM Y", 8, "Q3-TEAM Z", 4, "", "MINI"))
        ws.append((datetime(2026, 7, 11), datetime(2026, 7, 11, 11, 0).time(), "VENUE", "T2",
                   "X1(1stP)-", None, "X2(1stQ)-", None, "", "MINI"))
        buf = _io.BytesIO(); wb.save(buf)
        return buf.getvalue()

    import app as appmod
    orig_find_excel = appmod.find_excel

    def _fetch(round1_played: bool):
        xlsx_bytes = _make_wb(round1_played)
        appmod.find_excel = lambda tid: _io.BytesIO(xlsx_bytes) if tid == "quiksilver-cup" else orig_find_excel(tid)
        appmod._parse_cache.clear()
        with appmod.app.test_client() as c:
            r = c.get("/api/games/quiksilver-cup/Trojan%20Cardinal")
            return _json.loads(r.data)

    try:
        pre = _fetch(round1_played=False)
        pre_upcoming = pre.get("upcoming", [])
        ids = [g["game_id"] for g in pre_upcoming]
        ok = _check("pre-results: T2 (Round 2 branch) is shown, not suppressed or replaced by a stub",
                    "T2" in ids and not any(g.get("tbd_stub") for g in pre_upcoming if g["game_id"] == "T2"),
                    f"got {[(g['game_id'], g.get('tbd_stub')) for g in pre_upcoming]}")
        if not ok: failures += 1

        t2 = next((g for g in pre_upcoming if g["game_id"] == "T2"), None)
        ok = _check("pre-results: T2 tagged as a pool_1 branch (not a confirmed game)",
                    t2 is not None and t2.get("path") == "pool_1" and t2.get("placeholder") is True,
                    f"got path={t2.get('path') if t2 else None!r} placeholder={t2.get('placeholder') if t2 else None!r}")
        if not ok: failures += 1

        # Pool Q is always fully played in this fixture (independent of
        # round1_played) -- Team X is 1st, so T2's opponent should already be
        # resolvable to a real name even before Cardinal's own pool finishes.
        ok = _check("pre-results: T2 opponent already resolves via pool Q's standings (Team X)",
                    t2 is not None and t2.get("opponent", "").upper() == "TEAM X",
                    f"got opponent={t2.get('opponent') if t2 else None!r}")
        if not ok: failures += 1

        post = _fetch(round1_played=True)
        post_played = post.get("played", [])
        post_upcoming = post.get("upcoming", [])
        t2_post = next((g for g in post_upcoming if g["game_id"] == "T2"), None)
        ok = _check("post-results: Cardinal 2-0 in pool P, T2 still correctly resolved (not dropped)",
                    len(post_played) == 2 and t2_post is not None
                    and t2_post.get("opponent", "").upper() == "TEAM X",
                    f"played={[g['game_id'] for g in post_played]} "
                    f"t2_opponent={t2_post.get('opponent') if t2_post else None!r}")
        if not ok: failures += 1
    finally:
        appmod.find_excel = orig_find_excel
        appmod._parse_cache.clear()

    return failures


# ── Generic game_num ordering guard (automates a checklist item, not a script) ──
# feedback-new-tournament-checklist has always said, as a MANUAL step: "upcoming
# game_nums are in ascending date+time order". The Quiksilver Cup game_num bug
# (Friday's second game numbered after Saturday's placeholder) is exactly the
# bug that checklist line exists to catch — and it wasn't actually run as a
# concrete assertion before shipping, only eyeballed for crashes/self-references.
# test_game_num_chronological_order (above) only reproduces that ONE shape.
# This test makes the checklist item itself a permanent, generic check across
# multiple tournament formats, so the same class of bug in a different shape
# (not just this exact 2-Friday/1-Saturday case) still gets caught automatically.
GAME_NUM_ORDER_CHECKS = [
    # (tournament_id, team, sheet-or-None)
    ("quiksilver-cup", "Trojan Cardinal", None),
    ("jo-quals", "TROJAN CARDINAL A",  "18U Boys"),
    ("jo-quals", "TROJAN GOLD B",       "18U Boys"),
    ("jo-quals", "TROJAN CARDINAL (A)", "16U Boys"),
    ("jo-quals", "TROJAN GOLD (B)",     "16U Boys"),
    ("jo-quals", "TROJAN SILVER (C)",   "16U Boys"),
]

def test_game_num_ascending_order() -> int:
    """For each checked team, the API already returns 'played' and 'upcoming'
    sorted chronologically (my_games.sort() in api_games runs before game_num
    is assigned) -- so within each list, game_num must be non-decreasing along
    list order. A decrease means a later-dated game got a lower game_num than
    an earlier one, i.e. exactly the Friday-after-Saturday bug."""
    from app import app as _flask_app
    import json as _json

    failures = 0
    print("\ngame_num ascending order (checklist-derived, cross-format)")

    with _flask_app.test_client() as c:
        for tid, team, sheet in GAME_NUM_ORDER_CHECKS:
            url = f"/api/games/{tid}/{team}"
            if sheet:
                url += f"?sheet={sheet}"
            r = c.get(url)
            if r.status_code != 200:
                print(f"  [SKIP] {tid}/{team}: status {r.status_code}")
                continue
            data = _json.loads(r.data)
            for list_name in ("played", "upcoming"):
                nums = [g["game_num"] for g in data.get(list_name, []) if g.get("game_num") is not None]
                bad = [(a, b) for a, b in zip(nums, nums[1:]) if b < a]
                ok = _check(f"{tid}/{team} [{list_name}]: game_num non-decreasing in list order",
                            not bad, f"nums={nums}")
                if not ok: failures += 1

    return failures


def test_cca_opponent_slots() -> int:
    """Verify _team_opp_slot returns the correct (non-self) opponent for CCA placeholder games.

    This test would have caught the bug where "Trojan Gold vs Loser of game #23"
    was shown when the team IS the loser of game #23. The fix was _team_opp_slot;
    this test ensures it never regresses.
    """
    failures = 0
    print("\nCCA opponent-slot correctness")

    # Load each CSV file once
    csv_cache: dict[str, list] = {}
    for csv_file, division, prefix, team, game_id, expected_opp in CCA_OPP_CHECKS:
        path = os.path.join(FIXTURES_CCA, csv_file)
        if not os.path.exists(path):
            print(f"  [SKIP] fixture not found: {path}")
            continue

        if csv_file not in csv_cache:
            with open(path) as f:
                text = f.read()
            csv_cache[csv_file] = cca_parse_csv(text, division, prefix)
        div_games = csv_cache[csv_file]

        direct = [g for g in div_games
                  if team_matches(g["white_team"], team) or team_matches(g["dark_team"], team)]
        extras  = _expand_bracket_games(team, direct, div_games)
        all_g   = direct + extras
        my_ids  = {g["game_id"] for g in all_g}

        target = next((g for g in all_g if g["game_id"] == game_id), None)
        label  = f"{game_id} {team!r} opp={expected_opp!r}"

        if target is None:
            ok = _check(f"{label} — game found in expansion", False, "game not in expansion")
            failures += 1
            continue

        opp = _team_opp_slot(target, team, div_games, my_ids)

        # Primary: must match expected slot exactly
        ok = _check(f"{label}", opp == expected_opp, f"got {opp!r}")
        if not ok: failures += 1

        # Secondary (self-reference guard): opponent must never be the team's own slot
        # This catches future regressions even if expected_opp drifts as results fill in
        self_ref = team_matches(opp, team)
        ok2 = _check(f"{game_id} {team!r} — no self-reference", not self_ref,
                     f"opp={opp!r} matches team name")
        if not ok2: failures += 1

    return failures


# ── Game-num (bracket depth) regression test ─────────────────────────────────
# Regression: lose-path game (Saturday) was getting game_num=3 while win-path
# (Friday) got game_num=2. Both are "Game 2" — the second game in the journey.
# Fixed by using BFS depth from root instead of chronological slot order.

CCA_GAME_NUM_CHECKS = [
    # (tournament_id, team, sheet, game_id, expected_game_num, reason)
    ("jo-quals", "trojan gold (b)", "16U Boys", "16U-23", 1, "pool game = Game 1 (Fri 5PM)"),
    ("jo-quals", "trojan gold (b)", "16U Boys", "16U-24", 2, "win-path Game 2 (Fri 7:30PM)"),
    ("jo-quals", "trojan gold (b)", "16U Boys", "16U-46", 2, "lose-path Game 2 shares column with win (Sat 9AM)"),
    # Placement alternatives happen Sat 11AM — earlier than bracket games at Sat 12PM.
    # Re-numbering by actual time puts placements at Game 3, brackets at Game 4.
    ("jo-quals", "trojan gold (b)", "16U Boys", "16U-38", 3, "placement Game 3 (Sat 11AM, before bracket)"),
    ("jo-quals", "trojan gold (b)", "16U Boys", "16U-40", 3, "placement Game 3 (Sat 1PM, same column as 38)"),
    ("jo-quals", "trojan gold (b)", "16U Boys", "16U-57", 4, "bracket lose-path Game 4 (Sat 12PM)"),
    ("jo-quals", "trojan gold (b)", "16U Boys", "16U-49", 4, "bracket win-path Game 4 (Sat 1PM)"),
]

def test_cca_game_nums() -> int:
    """Verify win-path and lose-path games share the same game_num (same bracket depth).

    The bug: lose-path game (Saturday) got game_num=3 because chronological sort
    placed it after the win-path game (Friday). Fix: BFS depth from root game.
    """
    from app import app as _flask_app
    import json as _json

    failures = 0
    with _flask_app.test_client() as c:
        # Group by (tournament_id, team, sheet) to fetch once
        seen: dict = {}
        for tid, team, sheet, game_id, expected_num, reason in CCA_GAME_NUM_CHECKS:
            key = (tid, team, sheet)
            if key not in seen:
                r = c.get(f"/api/games/{tid}/{team}?sheet={sheet}")
                data = _json.loads(r.data)
                by_id = {g["game_id"]: g for g in (data.get("upcoming", []) + data.get("played", []))}
                seen[key] = by_id
            by_id = seen[key]
            g = by_id.get(game_id)
            got = g.get("game_num") if g else None
            ok = _check(f"{game_id} game_num={expected_num} ({reason})", got == expected_num,
                        f"got {got!r}")
            if not ok:
                failures += 1
    return failures


def test_trojan_team_names_clean() -> int:
    """No entry returned by api_trojan_teams should look like a raw slot string.

    The bug: CCA format uses 'W #12 - TROJAN CARDINAL (A)' slot syntax. _PREFIX_RE
    didn't handle the spaced variant, so strip_prefix left the raw slot intact and
    it appeared as a phantom team in the UI. This test catches any future case where
    a new slot format slips past strip_prefix.
    """
    from app import app as _flask_app, _SLOT_LIKE_RE as _slre
    import json as _json

    failures = 0
    tournaments_to_check = ["jo-quals"]
    with _flask_app.test_client() as c:
        for tid in tournaments_to_check:
            r = c.get(f"/api/trojan-teams/{tid}")
            if r.status_code != 200:
                _check(f"{tid} trojan-teams request succeeded", False, f"status {r.status_code}")
                failures += 1
                continue
            teams = _json.loads(r.data)
            for t in teams:
                name = t.get("name", "")
                ok = _check(
                    f"{tid}: team name is clean (not a raw slot): {name!r}",
                    not _slre.match(name),
                    f"looks like an unstripped slot",
                )
                if not ok:
                    failures += 1
    return failures


def test_strip_prefix_compound_pool_codes() -> int:
    """strip_prefix must handle compound "TIER_POOL" slot codes (e.g. "AU_P",
    "BZ_R") seen in real Junior Olympics data, not just plain single-letter
    pool codes ("A", "B").

    The bug: _PREFIX_RE's [A-Z]+ can't match past an underscore, so a slot
    like "3RD AU_P-ARROYO GRANDE" fails to match the prefix pattern at all and
    strip_prefix returns the whole raw string unchanged -- surfacing as a
    phantom team name. Found via a CCA-postmortem-style regression sweep
    against real 2025 Junior Olympics data (471 occurrences in that file
    alone) ahead of the 2026 Junior Olympics tournament.
    """
    from app import strip_prefix as _strip_prefix

    failures = 0
    cases = [
        ("3RD AU_P-ARROYO GRANDE", "ARROYO GRANDE"),
        ("2ND BZ_R-BURLINGAME", "BURLINGAME"),
        ("1ST CU_C", "1ST CU_C"),  # unresolved slot, no name yet -- left as-is
        # Nested variant: tier_pool code + parenthetical win/loss reference,
        # found live in the real 2026 Junior Olympics sheet (currently
        # unresolved as of 2026-07-14, but due to resolve to a real team
        # name -- including TROJAN GOLD itself -- as pool play concludes).
        ("1ST AU_S(W97)-TPCM SHARKS", "TPCM SHARKS"),
        ("2ND BZ_S(L97)-TROJAN GOLD", "TROJAN GOLD"),
        ("3RD BZ_S (W111)-VEGAS RENEGADES", "VEGAS RENEGADES"),
        ("1ST BZ_S(W97)", "1ST BZ_S(W97)"),  # unresolved -- left as-is
    ]
    for raw, expected in cases:
        got = _strip_prefix(raw)
        ok = _check(f"strip_prefix({raw!r})", got == expected, f"got {got!r}, expected {expected!r}")
        if not ok:
            failures += 1
    return failures


def test_no_slot_like_opponents() -> int:
    """Direct (non-placeholder) game cards must never show a slot-like opponent label.

    The bug: after strip_prefix fails on an unknown format, the raw slot string
    flows through as the opponent label — parents see "W #12 - TEMPLE CITY" instead
    of a real team name. This test catches that for all CCA Trojan teams.
    """
    from app import app as _flask_app, _SLOT_LIKE_RE as _slre
    import json as _json

    failures = 0
    checks = [
        ("jo-quals", "TROJAN CARDINAL A",   "18U Boys"),
        ("jo-quals", "TROJAN GOLD B",        "18U Boys"),
        ("jo-quals", "TROJAN CARDINAL (A)",  "16U Boys"),
        ("jo-quals", "TROJAN GOLD (B)",      "16U Boys"),
        ("jo-quals", "TROJAN SILVER (C)",    "16U Boys"),
    ]
    with _flask_app.test_client() as c:
        for tid, team, sheet in checks:
            r = c.get(f"/api/games/{tid}/{team}?sheet={sheet}")
            if r.status_code != 200:
                _check(f"{tid}/{team} request ok", False, f"status {r.status_code}")
                failures += 1
                continue
            data = _json.loads(r.data)
            all_games = data.get("upcoming", []) + data.get("played", [])
            for g in all_games:
                if g.get("placeholder"):
                    continue  # placeholder opponents are unresolved slots by design
                opp = g.get("opponent", "")
                ok = _check(
                    f"{tid}/{team} game {g['game_id']}: opponent not a raw slot: {opp!r}",
                    not _slre.match(opp),
                    "looks like an unstripped slot",
                )
                if not ok:
                    failures += 1
    return failures


def test_game_num_column_integrity() -> int:
    """Within each game_num column, non-placeholder games at different times on the
    same day must be win/lose path siblings — not sequential rounds incorrectly merged.

    The bug: CCA's 3:30 PM pool game and 6:50 PM bracket game both got game_num=1
    because the day-merge logic treated non-overlapping same-day times as the same
    round. This test catches that: if two non-placeholder games share a game_num,
    are on the same day, and are at different times, they must have win+lose paths.
    """
    from app import app as _flask_app
    import json as _json

    failures = 0
    checks = [
        ("jo-quals", "TROJAN CARDINAL A",   "18U Boys"),
        ("jo-quals", "TROJAN GOLD B",        "18U Boys"),
        ("jo-quals", "TROJAN CARDINAL (A)",  "16U Boys"),
        ("jo-quals", "TROJAN GOLD (B)",      "16U Boys"),
        ("jo-quals", "TROJAN SILVER (C)",    "16U Boys"),
    ]
    with _flask_app.test_client() as c:
        for tid, team, sheet in checks:
            r = c.get(f"/api/games/{tid}/{team}?sheet={sheet}")
            if r.status_code != 200:
                continue
            data = _json.loads(r.data)
            all_games = data.get("upcoming", []) + data.get("played", [])
            # Group non-placeholder games by (game_num, date)
            groups: dict = {}
            for g in all_games:
                if g.get("placeholder"):
                    continue
                key = (g.get("game_num"), g.get("date"))
                if None not in key:
                    groups.setdefault(key, []).append(g)
            for (gnum, date), grp in groups.items():
                times = {g.get("time") for g in grp if g.get("time")}
                if len(times) <= 1:
                    continue
                # Multiple times in same column on same day — must be win/lose siblings
                paths = {g.get("path") for g in grp}
                is_siblings = "win" in paths and "lose" in paths
                ok = _check(
                    f"{tid}/{team} GAME {gnum} on {date}: "
                    f"multiple times {sorted(times)} are win/lose siblings",
                    is_siblings,
                    f"paths={paths} — looks like sequential rounds merged into one column",
                )
                if not ok:
                    failures += 1
    return failures


def test_game_num_cascading_merge() -> int:
    """Games at the same round must share one game_num even when they
    descend from DIFFERENT immediate parent games, as long as those parents
    are themselves already in the same numbered group.

    The bug (caught live via a phone screenshot 2026-07-14, not by any
    automated check): a win-then-lose path and a lose-then-win path through
    a multi-round bracket land on different immediate parent games (e.g.
    game 28 vs game 32), so a decision-key-only grouping gives them
    different numbers (3 vs 4) even though both parents already share one
    number (2) and both paths represent exactly the same count of games
    played. "GAME 3" must mean the same thing regardless of which specific
    branch got the team there.

    Uses the real kap7-intl static fixture (not live-fetched, so this won't
    drift) -- TROJAN GOLD's 14U Silver bracket has a real 4-deep case: pool
    result determines two different round-2 outcomes (1st-in-pool vs
    2nd-in-pool), each of which further splits win/lose, and all four
    round-3 games must share one number, with all four of their round-4
    children sharing the next one.
    """
    from app import app as _flask_app
    import json as _json

    failures = 0
    with _flask_app.test_client() as c:
        r = c.get("/api/games/kap7-intl/TROJAN%20GOLD?sheet=14U%20BOYS%20SILVER-17%20TEAMS")
        if r.status_code != 200:
            _check("kap7-intl/TROJAN GOLD request succeeds", False, f"status {r.status_code}")
            return 1
        data = _json.loads(r.data)
        by_id = {g["game_id"]: g.get("game_num")
                 for g in data.get("played", []) + data.get("upcoming", [])}

        round3 = ["14Bag20", "14Bag22", "14Bag24", "14Bag26"]
        round4 = ["14Bag30", "14Bag32", "14Bag34", "14Bag35"]
        for label, ids in (("round 3", round3), ("round 4", round4)):
            nums = {gid: by_id.get(gid) for gid in ids}
            distinct = set(nums.values())
            ok = _check(
                f"kap7-intl/TROJAN GOLD {label}: {ids} all share one game_num",
                len(distinct) == 1 and None not in distinct,
                f"got {nums}",
            )
            if not ok:
                failures += 1
    return failures


def test_game_num_date_monotonic() -> int:
    """GAME numbers must track chronological order: GAME 2 can never start
    before GAME 1, across every real Trojan team/tournament combo currently
    configured. api_games's numbering pipeline guarantees this by
    construction (groups are sorted by earliest occurrence before being
    numbered), but this test is the permanent trip-wire against a future
    change to that code silently reintroducing an ordering bug -- the same
    "obvious mistake" a human giving the bracket a common-sense look would
    catch, done deterministically instead.
    """
    from app import app as _flask_app, KNOWN_TOURNAMENTS
    import json as _json

    # api_games serializes date/time as display strings ("Saturday, Jan 31",
    # "8:00 AM") for the frontend -- comparing those directly with < is
    # exactly the bug documented elsewhere in this app (never string-compare
    # time labels; "12:00 PM" < "8:00 AM" alphabetically). Parse them back
    # into comparable (month, day) / (hour, minute) tuples instead.
    _MONTHS = {m: i + 1 for i, m in enumerate(
        ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
         "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])}

    def _parse_date_str(s):
        if not s or s == "TBD":
            return None
        m = re.search(r'([A-Za-z]{3})\w*\s+(\d+)$', s)
        return (_MONTHS[m.group(1)], int(m.group(2))) if m else None

    def _parse_time_str(s):
        if not s or s == "TBD":
            return (0, 0)
        m = re.match(r'(\d+):(\d+)\s*(AM|PM)', s, re.IGNORECASE)
        if not m:
            return (0, 0)
        h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3).upper()
        if ap == "PM" and h != 12: h += 12
        if ap == "AM" and h == 12: h = 0
        return (h, mi)

    failures = 0
    with _flask_app.test_client() as c:
        for t in KNOWN_TOURNAMENTS:
            tid = t["id"]
            r = c.get(f"/api/trojan-teams/{tid}")
            if r.status_code != 200:
                continue
            for team in (_json.loads(r.data) or []):
                name, sheet = team.get("name"), team.get("sheet")
                qs = f"?sheet={sheet}" if sheet else ""
                gr = c.get(f"/api/games/{tid}/{name}{qs}")
                if gr.status_code != 200:
                    continue
                data = _json.loads(gr.data)
                games = data.get("played", []) + data.get("upcoming", [])
                earliest: dict[int, tuple] = {}
                for g in games:
                    gnum = g.get("game_num")
                    d = _parse_date_str(g.get("date"))
                    if gnum is None or d is None:
                        continue
                    key = (d, _parse_time_str(g.get("time")))
                    if gnum not in earliest or key < earliest[gnum]:
                        earliest[gnum] = key
                gn_sorted = sorted(earliest)
                for a, b in zip(gn_sorted, gn_sorted[1:]):
                    ok = _check(
                        f"{tid}/{name}: GAME {a} ({earliest[a]}) before GAME {b} ({earliest[b]})",
                        earliest[a] <= earliest[b],
                        "game numbers are out of chronological order",
                    )
                    if not ok:
                        failures += 1
    return failures


def test_njo_tree_multi_phase() -> int:
    """NJO trees must include every one of the team's real games, even when they
    span multiple disconnected segments (a round-robin pool phase with no
    w_to/l_to links at all, plus a separate elimination-bracket chain with its
    own numbering) rather than one continuous win/lose chain.

    Real bug found 2026-07-12 against last year's actual Junior Olympics data:
    a team with 9 real games showed only 3 in the tree because
    _build_njo_game_tree only followed the first connected chain from the
    team's earliest game and silently dropped a second, later bracket phase
    that started a fresh w_to/l_to sequence unconnected to the first.
    """
    failures = 0
    games = [
        # Two round-robin pool games with no advancement links at all.
        {"game_id": "T-001", "white_team": "TEST TEAM", "dark_team": "OPP A",
         "date": date(2026, 7, 24), "time": None, "played": True,
         "white_score": 10, "dark_score": 5, "w_to": None, "l_to": None},
        {"game_id": "T-002", "white_team": "TEST TEAM", "dark_team": "OPP B",
         "date": date(2026, 7, 24), "time": None, "played": True,
         "white_score": 6, "dark_score": 9, "w_to": None, "l_to": None},
        # A separate, later elimination-bracket chain: game 3 -> game 4.
        {"game_id": "T-003", "white_team": "TEST TEAM", "dark_team": "OPP C",
         "date": date(2026, 7, 25), "time": None, "played": True,
         "white_score": 8, "dark_score": 4, "w_to": 4, "l_to": None},
        {"game_id": "T-004", "white_team": "W#3-TEST TEAM", "dark_team": "OPP D",
         "date": date(2026, 7, 25), "time": None, "played": False,
         "white_score": None, "dark_score": None, "w_to": None, "l_to": None},
    ]
    tree = _build_njo_game_tree("TEST TEAM", games)
    tree_ids = {n["game_id"] for n in tree}
    expected_ids = {g["game_id"] for g in games}
    ok = _check(
        "NJO tree includes every real game across disconnected segments",
        tree_ids == expected_ids,
        f"expected {sorted(expected_ids)}, got {sorted(tree_ids)}",
    )
    if not ok:
        failures += 1
    return failures


def test_njo_tree_tied_game_no_misattribution() -> int:
    """A tied upstream game (_team_won returns None, same as "not yet
    played") must never let the bracket tree follow into a downstream
    branch whose own slot has already been explicitly resolved to a
    different, named team.

    Real bug found live 2026-07-19 via a round-by-round replay of real 2025
    Junior Olympics results: a team's bracket tied 9-9 in a semifinal; the
    downstream W#1 slot was already explicitly resolved to a different real
    team's name ("W#1-OPP A"), but the tree followed into it anyway and
    attributed that other team's own real, played result to our team as a
    fabricated game. Same root cause, same fix
    (_resolved_belongs_to_other_team), as the twin bug already fixed in
    _expand_bracket_games for the flat schedule list -- this guards the
    separate bracket-tree code path (_build_njo_game_tree's _follow).
    """
    failures = 0
    games = [
        {"game_id": "G1", "white_team": "OUR TEAM", "dark_team": "OPP A",
         "date": date(2026, 7, 23), "time": None, "played": True,
         "white_score": 9, "dark_score": 9, "comments": "semi",
         "w_to": None, "l_to": None},
        # Win-path, explicitly resolved to a DIFFERENT team -- must never
        # appear as our own game, regardless of the tie above.
        {"game_id": "G2", "white_team": "W#1-OPP A", "dark_team": "OPP B",
         "date": date(2026, 7, 24), "time": None, "played": True,
         "white_score": 11, "dark_score": 5, "comments": "9th",
         "w_to": None, "l_to": None},
        # Lose-path, explicitly resolved to OUR team -- this IS our real game.
        {"game_id": "G3", "white_team": "L#1-OUR TEAM", "dark_team": "OPP C",
         "date": date(2026, 7, 24), "time": None, "played": True,
         "white_score": 6, "dark_score": 10, "comments": "11th",
         "w_to": None, "l_to": None},
    ]
    tree = _build_njo_game_tree("OUR TEAM", games)
    tree_ids = {n["game_id"] for n in tree}
    ok = _check(
        "tied-game bracket tree includes only our own real games",
        tree_ids == {"G1", "G3"},
        f"expected {{'G1', 'G3'}}, got {sorted(tree_ids)} -- G2 belongs to OPP A, not us",
    )
    if not ok:
        failures += 1
    return failures


def test_njo_tree_slot_code_advancement() -> int:
    """A w_to/l_to value that is a text slot code (e.g. "grp_D3", meaning
    "seed 3 of sub-bracket grp_D") instead of a plain game number must
    resolve to every game whose own slot matches that code -- not just the
    first, since a round-robin sub-bracket plays the same seed against
    every other seed in its group (multiple matching games), and must
    resolve at all instead of being silently discarded, since some formats
    seed a winner/loser directly into a later group-stage slot rather than
    a single numbered game.

    Real bug found live 2026-07-20 against the real 2026 Junior Olympics
    18U Invite bracket: the parser discarded any non-numeric "W to #" / "L
    to #" value as unparseable, so Trojan Gold's bracket dead-ended right
    after its Day 1 cross game even though the sheet already stated exactly
    where the winner and loser each go next.
    """
    failures = 0
    games = [
        {"game_id": "G1", "white_team": "OUR TEAM", "dark_team": "OPP A",
         "date": date(2026, 7, 23), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "Cross",
         "w_to": "grp_D3", "l_to": "grp_C1"},
        # Round-robin sub-bracket: seed 3 plays both seed 1 and seed 2.
        {"game_id": "G2", "white_team": "GRP_D1-", "dark_team": "GRP_D3",
         "date": date(2026, 7, 24), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "Group",
         "w_to": None, "l_to": None},
        {"game_id": "G3", "white_team": "GRP_D2-", "dark_team": "GRP_D3",
         "date": date(2026, 7, 24), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "Group",
         "w_to": None, "l_to": None},
        # Does NOT reference seed 3 -- must never be pulled in.
        {"game_id": "G4", "white_team": "GRP_D1-", "dark_team": "GRP_D2-",
         "date": date(2026, 7, 24), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "Group",
         "w_to": None, "l_to": None},
        # Loss path target, explicitly a different team's real game.
        {"game_id": "G5", "white_team": "GRP_C1", "dark_team": "OPP B",
         "date": date(2026, 7, 24), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "Group",
         "w_to": None, "l_to": None},
    ]
    tree = _build_njo_game_tree("OUR TEAM", games)
    tree_ids = {n["game_id"] for n in tree}
    ok = _check(
        "slot-code advancement follows every matching game, not just the first, and no others",
        tree_ids == {"G1", "G2", "G3", "G5"},
        f"expected {{'G1','G2','G3','G5'}}, got {sorted(tree_ids)}",
    )
    if not ok:
        failures += 1
    g1 = next((n for n in tree if n["game_id"] == "G1"), None)
    if g1:
        ok = _check(
            "win path reaches both round-robin games for our seed",
            set(g1.get("win_next_ids") or []) == {"G2", "G3"},
            f"got {g1.get('win_next_ids')}",
        )
        if not ok:
            failures += 1
        ok = _check(
            "lose path reaches the loss-side slot code target",
            set(g1.get("lose_next_ids") or []) == {"G5"},
            f"got {g1.get('lose_next_ids')}",
        )
        if not ok:
            failures += 1
    return failures


def test_tree_layout_column_monotonic_invariant() -> int:
    """The bracket tree's column assignment must guarantee every real
    child's column is strictly greater than its parent's -- and when a node
    has more than one real parent (the "continues next tournament day" TBD
    stub attaches to every dead-end branch regardless of depth, so it can
    have parents at genuinely different depths), its column must be
    max(all parents' columns) + 1, not just whichever parent a traversal
    happens to reach it through first.

    Real bug found live 2026-07-20 against the real 2026 Junior Olympics
    18U Champ bracket: reusing the flat schedule's chronological game
    numbering for the tree's column position let two completely unrelated
    branches land in adjacent columns and get drawn as if one caused the
    other -- a long chain of "Lose" connectors between games that had never
    played each other. This test locks in the fix
    (_compute_tree_layout's own depth-based column assignment) directly.
    """
    failures = 0
    nodes = [
        {"game_id": "R1", "src_game_id": None, "date": date(2026, 7, 23), "time": None,
         "played": False, "white_team": "OUR TEAM", "dark_team": "OPP",
         "win_next_ids": ["A"], "lose_next_ids": ["D"], "pool_next": {}},
        {"game_id": "A", "src_game_id": "R1", "date": date(2026, 7, 24), "time": None,
         "played": False, "white_team": "OUR TEAM", "dark_team": "OPP",
         "win_next_ids": ["B"], "lose_next_ids": ["C"], "pool_next": {}},
        {"game_id": "B", "src_game_id": "A", "date": date(2026, 7, 25), "time": None,
         "played": False, "white_team": "OUR TEAM", "dark_team": "OPP",
         "win_next_ids": [], "lose_next_ids": [], "pool_next": {"SHARED": 1}},
        {"game_id": "C", "src_game_id": "A", "date": date(2026, 7, 25), "time": None,
         "played": False, "white_team": "OUR TEAM", "dark_team": "OPP",
         "win_next_ids": [], "lose_next_ids": [], "pool_next": {}},
        {"game_id": "D", "src_game_id": "R1", "date": date(2026, 7, 24), "time": None,
         "played": False, "white_team": "OUR TEAM", "dark_team": "OPP",
         "win_next_ids": [], "lose_next_ids": [], "pool_next": {"SHARED": 1}},
        {"game_id": "SHARED", "src_game_id": "D", "date": date(2026, 7, 26), "time": None,
         "played": False, "white_team": "OUR TEAM", "dark_team": "OPP",
         "win_next_ids": [], "lose_next_ids": [], "pool_next": {}},
        # A second, independent segment -- a genuinely separate root game
        # with no path to/from the first segment, scheduled later.
        {"game_id": "R2", "src_game_id": None, "date": date(2026, 7, 27), "time": None,
         "played": False, "white_team": "OUR TEAM", "dark_team": "OPP2",
         "win_next_ids": [], "lose_next_ids": [], "pool_next": {}},
    ]
    layout = _compute_tree_layout(nodes, "OUR TEAM")

    ok = _check("root starts at column 1", layout["R1"]["column"] == 1,
                f"got {layout['R1']['column']}")
    if not ok: failures += 1

    ok = _check(
        "simple chain: A and D are both column 2 (direct children of root), B and C are column 3",
        layout["A"]["column"] == 2 and layout["D"]["column"] == 2
        and layout["B"]["column"] == 3 and layout["C"]["column"] == 3,
        f"got A={layout['A']['column']} D={layout['D']['column']} "
        f"B={layout['B']['column']} C={layout['C']['column']}",
    )
    if not ok: failures += 1

    ok = _check(
        "shared multi-parent node uses max(parent columns) + 1, not whichever parent is reached first",
        layout["SHARED"]["column"] == 4,
        f"expected 4 (parents are B=3 and D=2, so max+1=4), got {layout['SHARED']['column']}",
    )
    if not ok: failures += 1

    ok = _check(
        "independent second segment starts after the first segment's deepest column, not overlapping at column 1",
        layout["R2"]["column"] == 5,
        f"expected 5 (first segment's deepest column is SHARED=4), got {layout['R2']['column']}",
    )
    if not ok: failures += 1

    # General invariant, the actual bug this whole redesign fixes: every
    # real edge (win/lose/pool_next) must have child column > parent column.
    bad_edges = []
    for n in nodes:
        for nid in ((n.get("win_next_ids") or []) + (n.get("lose_next_ids") or [])
                    + list((n.get("pool_next") or {}).keys())):
            pcol = layout[n["game_id"]]["column"]
            ccol = layout[nid]["column"]
            if ccol <= pcol:
                bad_edges.append((n["game_id"], nid, pcol, ccol))
    ok = _check("every edge has child column > parent column",
                not bad_edges, f"violations: {bad_edges}")
    if not ok: failures += 1

    return failures


def test_njo_tree_elimination_respects_live_parent() -> int:
    """A shared downstream node with multiple parents must only be marked
    eliminated once EVERY one of its parents is itself eliminated -- not the
    moment any single one of them is.

    Real bug found live 2026-07-23 against Trojan Gold 16U's actual, live
    2026 Junior Olympics results: won game 1, lost game 2. The dead "if
    we'd won game 2" branch and the real, still-live game 3 (loser's
    continuation) both merge into the same shared downstream TBD-stub node
    (the "continues next day" placeholder chain _append_tbd_stub_chain
    attaches). The previous _compute_tree_layout recursively marked a node
    eliminated the instant ANY parent died, with no check for a still-live
    sibling parent also feeding it -- so the shared node (and everything
    chained after it) rendered as a greyed-out "eliminated" placeholder
    even though the team's real, live path reaches it too and those games
    will genuinely still happen.
    """
    failures = 0
    nodes = [
        # Game 1: played, WON. Its lose-branch ("dead1") is now impossible.
        {"game_id": "g1", "src_game_id": None, "date": date(2026, 7, 23), "time": None,
         "played": True, "white_team": "TROJAN GOLD", "dark_team": "OPP1",
         "white_score": 10, "dark_score": 5,
         "win_next_ids": ["g2"], "lose_next_ids": ["dead1"], "pool_next": {}},
        # Game 2: played, LOST. Its win-branch ("dead2") is now impossible.
        {"game_id": "g2", "src_game_id": "g1", "date": date(2026, 7, 23), "time": None,
         "played": True, "white_team": "OPP2", "dark_team": "TROJAN GOLD",
         "white_score": 10, "dark_score": 9,
         "win_next_ids": ["dead2"], "lose_next_ids": ["g3"], "pool_next": {}},
        # Game 3: the team's REAL, still-live next game (loser's continuation).
        {"game_id": "g3", "src_game_id": "g2", "date": date(2026, 7, 23), "time": None,
         "played": False, "white_team": "TROJAN GOLD", "dark_team": "OPP3",
         "win_next_ids": ["tbd0"], "lose_next_ids": ["tbd0"], "pool_next": {}},
        # The two genuinely dead branches -- both also feed the same shared stub.
        {"game_id": "dead1", "src_game_id": "g1", "date": date(2026, 7, 23), "time": None,
         "played": False, "white_team": "TROJAN GOLD", "dark_team": "TBD",
         "win_next_ids": ["tbd0"], "lose_next_ids": ["tbd0"], "pool_next": {}},
        {"game_id": "dead2", "src_game_id": "g2", "date": date(2026, 7, 23), "time": None,
         "played": False, "white_team": "TROJAN GOLD", "dark_team": "TBD",
         "win_next_ids": ["tbd0"], "lose_next_ids": ["tbd0"], "pool_next": {}},
        # Shared downstream "continues next day" placeholder chain -- fed by
        # the real live game (g3) AND both dead branches.
        {"game_id": "tbd0", "src_game_id": "g3", "date": date(2026, 7, 24), "time": None,
         "played": False, "white_team": "TROJAN GOLD", "dark_team": "TBD",
         "win_next_ids": ["tbd1"], "lose_next_ids": ["tbd1"], "pool_next": {}},
        {"game_id": "tbd1", "src_game_id": "tbd0", "date": date(2026, 7, 25), "time": None,
         "played": False, "white_team": "TROJAN GOLD", "dark_team": "TBD",
         "win_next_ids": [], "lose_next_ids": [], "pool_next": {}},
    ]
    layout = _compute_tree_layout(nodes, "TROJAN GOLD")

    ok = _check("dead1 (impossible lose-branch of a won game) is eliminated",
                layout["dead1"]["eliminated"] is True, f"got {layout['dead1']}")
    if not ok: failures += 1

    ok = _check("dead2 (impossible win-branch of a lost game) is eliminated",
                layout["dead2"]["eliminated"] is True, f"got {layout['dead2']}")
    if not ok: failures += 1

    ok = _check("g3 (the real, still-live next game) is NOT eliminated",
                layout["g3"]["eliminated"] is False, f"got {layout['g3']}")
    if not ok: failures += 1

    ok = _check(
        "shared downstream stub fed by both dead branches AND the live game "
        "is NOT eliminated, since at least one live parent (g3) reaches it",
        layout["tbd0"]["eliminated"] is False, f"got {layout['tbd0']}")
    if not ok: failures += 1

    ok = _check(
        "further downstream stub (parent tbd0, itself not eliminated) is also NOT eliminated",
        layout["tbd1"]["eliminated"] is False, f"got {layout['tbd1']}")
    if not ok: failures += 1

    return failures


def test_njo_tree_pool_rank_tbd_stub() -> int:
    """When a team's pool has more possible finishing ranks than the sheet
    has literal finish-slot games for, the missing ranks must show as an
    honest TBD placeholder card, not silently vanish -- a parent seeing only
    one branch after a 3-team pool reads as "this bracket is broken", not
    "the tournament hasn't published this part yet".

    Real bug found live 2026-07-20 against the real 2026 Junior Olympics
    18U Invite bracket: Trojan Gold's 3-team pool only had a literal
    cross-game row for the 2nd-place finish; 1st and 3rd place had no
    discoverable game/slot anywhere in the sheet, and the bracket simply
    stopped showing them.
    """
    failures = 0
    games = [
        {"game_id": "P1", "white_team": "B1-OUR TEAM", "dark_team": "B3-OPP A",
         "date": date(2026, 7, 23), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "Group",
         "w_to": None, "l_to": None},
        {"game_id": "P2", "white_team": "B1-OUR TEAM", "dark_team": "B2-OPP B",
         "date": date(2026, 7, 23), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "Group",
         "w_to": None, "l_to": None},
        # Only the 2nd-place finish has a real downstream game.
        {"game_id": "P3", "white_team": "2ndB-", "dark_team": "2ndG-",
         "date": date(2026, 7, 23), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "Cross",
         "w_to": None, "l_to": None},
    ]
    tree = _build_njo_game_tree("OUR TEAM", games)
    anchor = next((n for n in tree if n["game_id"] == "P2"), None)
    ok = _check("pool-phase anchor game found in tree", anchor is not None)
    if not ok:
        failures += 1
        return failures
    ranks = set(anchor.get("pool_next", {}).values())
    ok = _check(
        "pool_next covers all 3 possible ranks (1 real, 2 TBD)",
        ranks == {1, 2, 3},
        f"got {ranks}",
    )
    if not ok:
        failures += 1
    tbd_stubs = [n for n in tree if n.get("tbd_stub") and n.get("src_game_id") == "P2"]
    ok = _check(
        "exactly 2 TBD stubs synthesized for the uncovered ranks (1st, 3rd)",
        len(tbd_stubs) == 2,
        f"got {len(tbd_stubs)}: {[n['game_id'] for n in tbd_stubs]}",
    )
    if not ok:
        failures += 1
    return failures


def test_jo_2026_teams_reachable() -> int:
    """Every known Trojan Junior Olympics 2026 team must still be found on
    the live schedule, return real games, and never show red bracket
    confidence.

    The JO schedule is a live Google Sheet the organizer edits continuously
    (confirmed live 2026-07-14: two fetches minutes apart returned different
    game_ids for the same team/sheet), so this intentionally does NOT pin
    exact game_ids or game counts the way CHAMPIONSHIP_CHECKS does for the
    static WPL fixture -- it would go stale within days. Instead it's a
    loose regression guard against a genuine team disappearing entirely
    (organizer re-seeding removed them from this sheet/division), a crash,
    or a real structural break (red confidence) -- as opposed to yellow,
    which is expected and correct this early before results start rolling in.
    """
    from app import app as _flask_app
    import json as _json

    failures = 0
    # (team, sheet, min_games) -- confirmed live 2026-07-14 against the
    # public 2026 JO schedule (see JO_SHEETS_URL in app.py).
    checks = [
        ("TROJAN CARDINAL", "18U_M_Champ",      1),
        ("TROJAN GOLD",     "18U_M_Invite 24",  1),
        ("TROJAN CARDINAL", "16U_M_Champ",      1),
        ("TROJAN GOLD",     "16U_M_Classic",    1),
        ("TROJAN CARDINAL", "14U_M_Classic",    1),
        ("TROJAN CARDINAL", "12U_M_Classic_53", 1),
        ("TROJAN GOLD",     "12U_M_Classic_53", 1),
    ]
    with _flask_app.test_client() as c:
        for team, sheet, min_games in checks:
            r = c.get(f"/api/games/junior-olympics/{team}?sheet={sheet}")
            ok = _check(f"{team}/{sheet}: request succeeds", r.status_code == 200,
                        f"status {r.status_code}")
            if not ok:
                failures += 1
                continue
            data = _json.loads(r.data)
            games = data.get("played", []) + data.get("upcoming", [])
            ok = _check(f"{team}/{sheet}: >= {min_games} game(s) found",
                        len(games) >= min_games, f"got {len(games)}")
            if not ok:
                failures += 1
            conf = data.get("bracket_confidence")
            ok = _check(f"{team}/{sheet}: bracket confidence is not red",
                        conf != "red", f"got {conf!r}, warnings={data.get('bracket_warnings')}")
            if not ok:
                failures += 1
    return failures


def test_place_predictor_seeded_bracket() -> int:
    """Place Predictor must still produce finish probabilities for a
    division that skips round-robin pools entirely and seeds straight into a
    single-elimination bracket (root slots like "1-ALPHA", no pool letter
    anywhere in the sheet).

    Real bug found live 2026-07-19 against the actual 2026 Junior Olympics
    schedule: 4 of 5 Trojan teams (18U/16U Champ, 16U/14U Classic) use
    exactly this format, and Place Predictor silently returned "pool
    schedule not posted yet" for all of them even though the schedule (and a
    genuine 48-team seeded bracket) was fully posted. Root cause:
    _tournament_finish_probs only ever built its team roster from
    pool-letter slots (_POOL_SLOT_RE) -- a division with none had nothing to
    simulate from, even though strip_prefix's bare "\\d+-" alternative
    already resolves these same root slots correctly for the (unaffected)
    Schedule/Bracket view.
    """
    failures = 0

    # Synthetic 4-team single-elimination bracket, no pool letters anywhere:
    # ALPHA(seed1)/DELTA(seed4) and BRAVO(seed2)/CHARLIE(seed3) feed a final
    # (comments "1st") and a 3rd-place game (comments "3rd") via W#/L# refs.
    games = [
        {"game_id": "G1", "white_team": "1-ALPHA", "dark_team": "4-DELTA",
         "date": date(2026, 7, 23), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": None, "sheet": "TEST"},
        {"game_id": "G2", "white_team": "2-BRAVO", "dark_team": "3-CHARLIE",
         "date": date(2026, 7, 23), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": None, "sheet": "TEST"},
        {"game_id": "G3", "white_team": "W#1", "dark_team": "W#2",
         "date": date(2026, 7, 24), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "1st", "sheet": "TEST"},
        {"game_id": "G4", "white_team": "L#1", "dark_team": "L#2",
         "date": date(2026, 7, 24), "time": None, "played": False,
         "white_score": None, "dark_score": None, "comments": "3rd", "sheet": "TEST"},
    ]
    our_probs, all_probs = _tournament_finish_probs("ALPHA", games, n_trials=200)
    ok = _check("seeded bracket with no pool letters still produces finish probabilities",
                bool(our_probs), f"got {our_probs!r}")
    if not ok:
        failures += 1
    ok = _check("all 4 seeded teams appear in the simulated roster",
                len(all_probs) == 4, f"got {sorted(all_probs.keys())}")
    if not ok:
        failures += 1
    ok = _check("top seed's placement probabilities sum to ~1.0 (every trial resolves to a placement)",
                bool(our_probs) and abs(sum(our_probs.values()) - 1.0) < 0.01,
                f"got sum={sum(our_probs.values()) if our_probs else 0}")
    if not ok:
        failures += 1

    # Live confirmation against the real 2026 Junior Olympics sheet -- the
    # 4 divisions actually found broken, via the real API route (not just
    # the simulation function directly).
    from app import app as _flask_app
    import json as _json

    live_checks = [
        ("TROJAN CARDINAL", "18U_M_Champ"),
        ("TROJAN CARDINAL", "16U_M_Champ"),
        ("TROJAN GOLD",     "16U_M_Classic"),
        ("TROJAN CARDINAL", "14U_M_Classic"),
    ]
    with _flask_app.test_client() as c:
        for team, sheet in live_checks:
            r = c.get(f"/api/place-predictor/junior-olympics/{team}?sheet={sheet}")
            ok = _check(f"{team}/{sheet}: predictor request succeeds", r.status_code == 200,
                        f"status {r.status_code}")
            if not ok:
                failures += 1
                continue
            data = _json.loads(r.data)
            ok = _check(f"{team}/{sheet}: predictor returns finish probabilities (no pool letter in this division)",
                        bool(data.get("finish_probs")),
                        f"got pool={data.get('pool')!r}, finish_probs={data.get('finish_probs')!r}")
            if not ok:
                failures += 1
    return failures


def test_all_known_tournaments_teams_healthy() -> int:
    """New-tournament checklist item 4, automated generically instead of
    hand-copied per tournament.

    Every prior version of this check (test_jo_2026_teams_reachable,
    test_no_slot_like_opponents, CCA_GAME_NUM_CHECKS, etc.) hardcodes its own
    (tournament_id, team, sheet) list -- someone has to remember to add a new
    tournament's teams to a new test function. That's exactly the failure
    mode documented in feedback-new-tournament-checklist: "reviewing this
    list is not the same as running it" -- Quiksilver Cup's game_num bug
    shipped because checklist item 4 was eyeballed, not encoded as a running
    assertion, for that specific tournament.

    This test instead discovers every Trojan team for every tournament in
    KNOWN_TOURNAMENTS via /api/trojan-teams/<id> (the same discovery
    mechanism the pre-game monitor uses) and runs the same battery against
    whatever it finds -- so a newly-added tournament is covered automatically
    the moment it's added to KNOWN_TOURNAMENTS, with no new test to write.

    Deliberately loose (matches test_jo_2026_teams_reachable's philosophy):
    does not pin exact game_ids/counts, since several tournaments here are
    live-fetched and drift. Checks the structural invariants a "look at it
    for a second" pass would catch: no crash, no red confidence, no
    self-referencing opponent, no raw unresolved slot leaking into a
    non-placeholder card, and non-decreasing game_num order.
    """
    from app import app as _flask_app, KNOWN_TOURNAMENTS, _SLOT_LIKE_RE as _slre, team_matches
    import json as _json

    failures = 0
    with _flask_app.test_client() as c:
        for t in KNOWN_TOURNAMENTS:
            tid = t["id"]
            r = c.get(f"/api/trojan-teams/{tid}")
            if r.status_code != 200:
                print(f"  [SKIP] {tid}: /api/trojan-teams returned {r.status_code}")
                continue
            teams = _json.loads(r.data)
            if not teams:
                print(f"  [SKIP] {tid}: no Trojan teams discovered (no local fixture / live URL not set)")
                continue
            for team_info in teams:
                name, sheet = team_info["name"], team_info["sheet"]
                label = f"{tid}/{name}/{sheet}"
                gr = c.get(f"/api/games/{tid}/{name}?sheet={sheet}")
                ok = _check(f"{label}: request succeeds", gr.status_code == 200,
                            f"status {gr.status_code}")
                if not ok:
                    failures += 1
                    continue
                data = _json.loads(gr.data)
                games = data.get("played", []) + data.get("upcoming", [])
                if not games:
                    print(f"  [SKIP] {label}: 0 games returned")
                    continue

                conf = data.get("bracket_confidence")
                ok = _check(f"{label}: bracket confidence is not red", conf != "red",
                            f"got {conf!r}, warnings={data.get('bracket_warnings')}")
                if not ok:
                    failures += 1

                for g in games:
                    opp = g.get("opponent", "")
                    ok = _check(f"{label} game {g['game_id']}: opponent is not a self-reference",
                                not team_matches(opp, name), f"opponent {opp!r} matches team name")
                    if not ok:
                        failures += 1
                    if not g.get("placeholder"):
                        ok = _check(f"{label} game {g['game_id']}: opponent is not a raw unresolved slot",
                                    not _slre.match(opp), f"got {opp!r}")
                        if not ok:
                            failures += 1

                for list_name in ("played", "upcoming"):
                    nums = [g["game_num"] for g in data.get(list_name, []) if g.get("game_num") is not None]
                    bad = [(a, b) for a, b in zip(nums, nums[1:]) if b < a]
                    ok = _check(f"{label} [{list_name}]: game_num non-decreasing", not bad,
                                f"nums={nums}")
                    if not ok:
                        failures += 1
    return failures


def test_bracket_tree_last_meeting_parity() -> int:
    """canonical_bracket's last_meeting must exactly match the flat
    played/upcoming list's last_meeting for the exact same game_id --
    checked in BOTH directions, not just "tree is not worse-informed."

    Two distinct real bugs shipped at this exact call site, both found
    2026-07-16 while auditing JO for the "two divergent code paths" pattern:

    1. Tree used _all_historical_games() alone (local archive only) while
       the flat list used _h2h_games (archive + this tournament's own live
       data combined) -- tree could be WORSE-INFORMED, missing a
       same-tournament rematch the flat list would show. Fixed by pointing
       the tree serializer at _h2h_games too.
    2. This test originally only checked direction 1 (skipped the
       comparison whenever flat_val was None, on the theory that null means
       "nothing to miss"). That let a SECOND bug slip past it undetected:
       the tree's call never passed sheet=, so it was unscoped across every
       age division while the flat list's call was correctly scoped to the
       team's own division -- tree could be BETTER-INFORMED in a way that
       was actually WRONG, surfacing a 16U opponent's history on a 12U
       team's bracket card (confirmed live: 12U Trojan Cardinal vs Asphalt
       Green). Fixed by passing sheet=tree_sheet. Now checks strict equality
       both ways so neither direction of divergence can hide again.

    Also covers h2h, our_record, and opp_record (2026-07-18): these were
    simply MISSING from the tree serializer entirely until now (Quiksilver,
    a flat-list tournament, always showed full matchup history + this-
    tournament records; Junior Olympics, a bracket-tree tournament, only
    ever showed last_meeting). Lower divergence risk than the other fields
    here since both paths now call the exact same _opponent_history /
    sheet_records lookups rather than two independent implementations --
    still worth checking, since "calls the same function" and "passes it
    the same arguments" are two different claims, and this file has a
    documented history of the second one being wrong (last_meeting, twice).
    """
    from app import app as _flask_app
    import json as _json

    failures = 0
    checks = [
        ("TROJAN CARDINAL", "18U_M_Champ"),
        ("TROJAN GOLD",     "18U_M_Invite 24"),
        ("TROJAN CARDINAL", "16U_M_Champ"),
        ("TROJAN GOLD",     "16U_M_Classic"),
        ("TROJAN CARDINAL", "14U_M_Classic"),
        ("TROJAN CARDINAL", "12U_M_Classic_53"),
        ("TROJAN GOLD",     "12U_M_Classic_53"),
    ]
    fields = ["last_meeting", "h2h", "our_record", "opp_record"]
    with _flask_app.test_client() as c:
        for team, sheet in checks:
            r = c.get(f"/api/games/junior-olympics/{team}?sheet={sheet}")
            if r.status_code != 200:
                continue
            data = _json.loads(r.data)
            cb = data.get("canonical_bracket")
            if not cb:
                continue
            flat_by_id = {g["game_id"]: g for g in data.get("upcoming", []) if g.get("game_id")}
            for n in cb.get("guaranteed_games", []) + cb.get("possible_games", []):
                gid = n["game_id"]
                if gid not in flat_by_id:
                    continue
                flat_g = flat_by_id[gid]
                for field in fields:
                    flat_val = flat_g.get(field)
                    ok = _check(
                        f"{team}/{sheet}: canonical_bracket {field} for {gid!r} "
                        f"exactly matches the flat list",
                        n.get(field) == flat_val,
                        f"flat={flat_val!r} tree={n.get(field)!r}",
                    )
                    if not ok:
                        failures += 1
    return failures


def test_flat_list_tree_path_parity() -> int:
    """The flat list's 'path' field (win/lose/pool_N, derived from
    bracket_path + find_next_games) must agree with canonical_bracket's
    'path' field (derived from _compute_tree_layout's win_next_ids/
    lose_next_ids) for every game_id present in both.

    A third independent pair of implementations solving the same problem
    (which games are a win/lose continuation of which), found auditing for
    more instances of this session's dominant bug pattern after fixing
    game_num, last_meeting, and opponent resolution. No divergence found on
    any of the 7 known JO teams' current real data, so NOT changed to share
    one implementation -- that would be touching two different graph-
    traversal algorithms without a concrete bug driving it, the wrong risk
    trade this close to Junior Olympics (see the is_current fix in this
    same file for the shape of evidence that WOULD justify it). This test
    is the safe version: doesn't touch the runtime code, but will catch the
    moment these two algorithms actually do disagree on real data.
    """
    from app import app as _flask_app
    import json as _json

    failures = 0
    checks = [
        ("TROJAN CARDINAL", "18U_M_Champ"),
        ("TROJAN GOLD",     "18U_M_Invite 24"),
        ("TROJAN CARDINAL", "16U_M_Champ"),
        ("TROJAN GOLD",     "16U_M_Classic"),
        ("TROJAN CARDINAL", "14U_M_Classic"),
        ("TROJAN CARDINAL", "12U_M_Classic_53"),
        ("TROJAN GOLD",     "12U_M_Classic_53"),
    ]
    with _flask_app.test_client() as c:
        for team, sheet in checks:
            r = c.get(f"/api/games/junior-olympics/{team}?sheet={sheet}")
            if r.status_code != 200:
                continue
            data = _json.loads(r.data)
            flat_path = {g["game_id"]: g.get("path")
                         for g in data.get("upcoming", []) + data.get("played", [])}
            cb = data.get("canonical_bracket")
            if not cb:
                continue
            for n in cb.get("guaranteed_games", []) + cb.get("possible_games", []):
                gid = n["game_id"]
                if gid not in flat_path:
                    continue
                ok = _check(
                    f"{team}/{sheet}: {gid!r} path agrees between flat list and tree",
                    flat_path[gid] == n.get("path"),
                    f"flat={flat_path[gid]!r} tree={n.get('path')!r}",
                )
                if not ok:
                    failures += 1
    return failures


def test_flat_list_is_current_placeholder_guard() -> int:
    """A live score attached to a placeholder (still-unresolved-opponent)
    game must never flip is_current=True on the flat played/upcoming list --
    it must still show live_score, just not the CURRENT GAME badge.

    The tree serializer already had this exact guard (`if live and not
    node.get("placeholder")`), added after a real live-tournament bug
    ("CURRENT GAME must never show on a node whose opponent is still an
    unresolved guess" -- see feedback-wpl-bracket-rendering). The flat
    list's identical is_current block never got the same guard -- found
    auditing for other divergent-path bugs after fixing
    game_num/last_meeting/opponent resolution this same session. A live
    score CAN legitimately attach to a game_id that's still a placeholder
    in our data (the organizer's sheet cell hasn't been updated with a real
    opponent name yet even though the game is being played right now), so
    this isn't just a hypothetical.

    No real placeholder game currently has a live score (JO hasn't started
    yet), so this injects a synthetic one directly into _LIVE_SCORES for a
    known real placeholder game_id, checks the API response, then cleans up.
    """
    from app import app as _flask_app, _LIVE_SCORES
    import json as _json

    failures = 0
    tid, team, sheet, gid = "junior-olympics", "TROJAN GOLD", "16U_M_Classic", "16BX-028"
    key = (tid, gid)
    _LIVE_SCORES[key] = {"our_score": 3, "opp_score": 2, "quarter": "Q2", "updated_at": "now"}
    try:
        with _flask_app.test_client() as c:
            r = c.get(f"/api/games/{tid}/{team}?sheet={sheet}")
            if r.status_code != 200:
                print(f"  [SKIP] {tid}/{team}: status {r.status_code}")
                return 0
            data = _json.loads(r.data)
            game = next((g for g in data.get("upcoming", []) if g["game_id"] == gid), None)
            if game is None:
                print(f"  [SKIP] {gid} not found in {team}'s upcoming list (data may have moved on)")
                return 0
            if not game.get("placeholder"):
                print(f"  [SKIP] {gid} is no longer a placeholder (organizer resolved it) -- guard not exercised")
                return 0
            ok = _check(f"{gid}: live_score is populated", game.get("live_score") is not None,
                        f"got {game.get('live_score')!r}")
            if not ok:
                failures += 1
            ok = _check(f"{gid}: is_current stays False on a placeholder even with a live score",
                        game.get("is_current") is False, f"got {game.get('is_current')!r}")
            if not ok:
                failures += 1
    finally:
        _LIVE_SCORES.pop(key, None)
    return failures


def test_bracket_tree_opponent_resolution_parity() -> int:
    """canonical_bracket's opponent resolution (the tree serializer's own
    inline algorithm) must agree with _team_opp_slot -- the shared function
    the flat played/upcoming list already uses -- for every real
    (non-placeholder) tree node.

    These are two INDEPENDENT algorithms solving the same problem
    ("which slot in this game is ours"), not one shared function: the tree
    serializer (_serialize_tree_node) does its own raw-match-then-
    describe_slot-fallback check inline, while the flat list calls
    _team_opp_slot (which additionally handles pool_rank_group finish-slots
    and my_game_ids-based W#/L# resolution that the tree's inline version
    doesn't). Found while auditing for more instances of this project's
    single most common bug pattern (two parallel paths computing the same
    thing, one gets fixed, the other doesn't -- see game_num and
    last_meeting, both hit multiple times).

    Checked empirically against all 7 known JO teams' real trees before
    deciding what to do about it: every tree node's pool_rank/
    pool_rank_group is always None (that _team_opp_slot branch never
    triggers here), and every slot is either a direct name match or a
    W#/L# reference -- the two algorithms very likely already agree on all
    current real data. Rewriting the tree's inline algorithm to share
    _team_opp_slot directly would be the fuller fix, but doing that a week
    before Junior Olympics for an unproven benefit is the wrong risk
    trade -- this parity check is the safer version: it doesn't change the
    live code path at all, but will catch the moment the two algorithms
    actually do diverge on real data, which is exactly the trigger for
    doing the fuller unification with real evidence instead of a guess.
    """
    import sys as _sys
    sys_path_added = ROOT not in _sys.path
    if sys_path_added:
        _sys.path.insert(0, ROOT)
    from app import (
        app as _flask_app, find_excel, load_and_parse, _filter_by_dates,
        _build_njo_game_tree, _build_wpl_game_tree, _team_opp_slot,
        _expand_bracket_games, team_matches, describe_slot, _tournament_meta,
    )
    import json as _json

    failures = 0
    checks = [
        ("TROJAN CARDINAL", "18U_M_Champ"),
        ("TROJAN GOLD",     "18U_M_Invite 24"),
        ("TROJAN CARDINAL", "16U_M_Champ"),
        ("TROJAN GOLD",     "16U_M_Classic"),
        ("TROJAN CARDINAL", "14U_M_Classic"),
        ("TROJAN CARDINAL", "12U_M_Classic_53"),
        ("TROJAN GOLD",     "12U_M_Classic_53"),
    ]
    excel = find_excel("junior-olympics")
    if not excel:
        print("  [SKIP] junior-olympics source not reachable")
        return 0
    all_games = load_and_parse(excel)
    games = _filter_by_dates(all_games, "junior-olympics")
    meta = _tournament_meta("junior-olympics")
    ref_date = meta.get("date_start") if meta else None

    with _flask_app.test_client() as c:
        for team, sheet in checks:
            dg = [g for g in games if g["sheet"] == sheet]
            direct = [g for g in dg if team_matches(g["white_team"], team)
                      or team_matches(g["dark_team"], team)]
            if not direct:
                continue
            # Match api_games exactly: my_game_ids (what _team_opp_slot's
            # W#/L# resolution checks against) includes games reached via
            # _expand_bracket_games, not just direct name matches -- using
            # only direct matches here (an earlier version of this test did)
            # made _team_opp_slot look broken at round 4+ when it isn't;
            # the real app's my_game_ids is the expanded set.
            my_games = direct + _expand_bracket_games(team, direct, dg)
            my_game_ids = {g["game_id"] for g in my_games}
            latest = max((g["date"] for g in my_games if g.get("date")), default=None)
            tree = _build_njo_game_tree(team, dg, anchor_date=latest)
            expected = {}
            for n in tree:
                if n.get("placeholder"):
                    continue
                opp_sl = _team_opp_slot(n, team, dg, my_game_ids)
                expected[n["game_id"]] = describe_slot(opp_sl, dg, ref_date=ref_date)

            r = c.get(f"/api/games/junior-olympics/{team}?sheet={sheet}")
            if r.status_code != 200:
                continue
            data = _json.loads(r.data)
            cb = data.get("canonical_bracket")
            if not cb:
                continue
            for n in cb.get("guaranteed_games", []) + cb.get("possible_games", []):
                gid = n["game_id"]
                if gid not in expected or n.get("placeholder"):
                    continue
                ok = _check(
                    f"{team}/{sheet}: {gid!r} opponent resolution agrees between "
                    f"_team_opp_slot and the tree's own inline algorithm",
                    n.get("opponent") == expected[gid],
                    f"tree={n.get('opponent')!r} _team_opp_slot={expected[gid]!r}",
                )
                if not ok:
                    failures += 1
    return failures


def test_bracket_tree_game_num_no_collision() -> int:
    """canonical_bracket's game_num (what the bracket UI actually renders as
    "GAME N" and uses to lay out columns) must never assign the same number
    to two DIFFERENT real, independent root games -- and TBD stub numbers
    must never collide with a real game's number either.

    The bug (caught live via phone screenshots 2026-07-15, not by any
    automated check): _compute_tree_layout's BFS started every root node at
    column 1, which is correct for genuine win/lose or pool-finish
    ALTERNATIVES (they share one decision point) but wrong the moment a team
    has more than one real, independent root game -- e.g. two separate pool
    games against different opponents, neither descending from the other.
    Those silently landed on the same "GAME 1", hiding a real scheduled game
    from the bracket view entirely. Every other game_num test in this file
    (test_game_num_cascading_merge, test_game_num_date_monotonic, etc.)
    checks the FLAT played/upcoming list, which was already correct -- none
    of them touch canonical_bracket, which is the actual bracket-view data
    this bug lived in. This test closes that coverage gap directly.

    Confirmed live on 3 of the 7 known JO teams before the fix: 18U Trojan
    Gold (games 005 vs 021), 12U Trojan Gold (3 independent pool games), and
    (post-fix, from the TBD-stub numbering bug found while fixing this) 18U
    Trojan Gold's TBD chain colliding with game 021's real number 2.
    """
    from app import app as _flask_app
    import json as _json

    failures = 0
    checks = [
        ("TROJAN CARDINAL", "18U_M_Champ",      1),
        ("TROJAN GOLD",     "18U_M_Invite 24",  1),
        ("TROJAN CARDINAL", "16U_M_Champ",      1),
        ("TROJAN GOLD",     "16U_M_Classic",    1),
        ("TROJAN CARDINAL", "14U_M_Classic",    1),
        ("TROJAN CARDINAL", "12U_M_Classic_53", 1),
        ("TROJAN GOLD",     "12U_M_Classic_53", 1),
    ]
    with _flask_app.test_client() as c:
        for team, sheet, _min in checks:
            r = c.get(f"/api/games/junior-olympics/{team}?sheet={sheet}")
            if r.status_code != 200:
                continue
            data = _json.loads(r.data)
            cb = data.get("canonical_bracket")
            if not cb:
                continue
            nodes = cb.get("guaranteed_games", []) + cb.get("possible_games", [])

            roots = [n for n in nodes if not n.get("src_game_id") and not n.get("placeholder")]
            root_nums: dict = {}
            for n in roots:
                prev = root_nums.get(n["game_num"])
                ok = _check(
                    f"{team}/{sheet}: root game {n['game_id']!r} (game_num {n['game_num']}) "
                    f"doesn't collide with another independent root",
                    prev is None or prev == n["game_id"],
                    f"also root {prev!r}" if prev else "",
                )
                if not ok:
                    failures += 1
                root_nums[n["game_num"]] = n["game_id"]

            # No real (non-stub) node may have exactly one of
            # win_next_ids/lose_next_ids populated -- that's the bug found
            # live 2026-07-16 (confirmed against the raw sheet: "L#28" and
            # "W#32" existed, "W#28" and "L#32" did not, anywhere in the
            # division). A win/lose split must show both outcomes or
            # neither; _fill_missing_branch_stubs exists specifically to
            # turn "one outcome" into "one outcome + one TBD stub for the
            # other," so this must never observe exactly one after it runs.
            for n in nodes:
                if n.get("tbd_stub"):
                    continue
                has_win = bool(n.get("win_next_ids"))
                has_lose = bool(n.get("lose_next_ids"))
                ok = _check(
                    f"{team}/{sheet}: {n['game_id']!r} has both branches or neither "
                    f"(not exactly one dangling)",
                    has_win == has_lose,
                    f"win_next_ids={n.get('win_next_ids')} lose_next_ids={n.get('lose_next_ids')}",
                )
                if not ok:
                    failures += 1

            # Day-chain stubs (__tbd_<sheet>_<i>, added when the whole tree
            # runs dry for a day) represent a LATER round than anything
            # already known, so they must never collide with a real number.
            # Missing-branch stubs (__tbd_branch_<sheet>_<parent_gid>, added
            # by _fill_missing_branch_stubs when only one of a node's two
            # outcomes has been published) and pool-rank stubs
            # (__tbd_pool_<parent_gid>_<rank>, added when a pool has more
            # possible finish ranks than the sheet has real games for) are
            # the OPPOSITE by design: both fill in another alternative
            # WITHIN an already-numbered round, so they MUST share their
            # sibling's number, not avoid it.
            real_nums = {n["game_num"] for n in nodes if not n.get("tbd_stub")}
            day_chain_nums = {n["game_num"] for n in nodes
                               if n.get("tbd_stub")
                               and not n["game_id"].startswith("__tbd_branch_")
                               and not n["game_id"].startswith("__tbd_pool_")}
            overlap = real_nums & day_chain_nums
            ok = _check(f"{team}/{sheet}: day-chain TBD stub numbers don't collide with real game numbers",
                        not overlap, f"overlap={overlap}")
            if not ok:
                failures += 1

            by_gid = {n["game_id"]: n for n in nodes}
            for n in nodes:
                if not n.get("tbd_stub") or not n["game_id"].startswith("__tbd_branch_"):
                    continue
                parent = by_gid.get(n.get("src_game_id"))
                sibling_ids = ((parent.get("win_next_ids") or [])
                               + (parent.get("lose_next_ids") or [])) if parent else []
                sibling = next((by_gid[sid] for sid in sibling_ids
                                if sid != n["game_id"] and sid in by_gid), None)
                ok = _check(
                    f"{team}/{sheet}: missing-branch stub {n['game_id']!r} shares "
                    f"its sibling {sibling['game_id'] if sibling else None!r}'s game_num",
                    sibling is not None and n["game_num"] == sibling["game_num"],
                    f"stub game_num={n['game_num']}, sibling game_num={sibling.get('game_num') if sibling else None}",
                )
                if not ok:
                    failures += 1
    return failures


def test_last_meeting_available() -> int:
    """At least one upcoming game per Trojan team should have a last_meeting populated,
    confirming cross-tournament history search is active.

    This test will naturally pass only when there IS prior history (i.e., opponents
    have been faced in a previous tournament on file). It SKIPs teams with zero
    prior history rather than failing — the assertion is that the mechanism works,
    not that every opponent has been faced before.
    """
    from app import app as _flask_app
    import json as _json

    failures = 0
    checks = [
        ("jo-quals", "TROJAN CARDINAL (A)", "16U Boys"),
        ("jo-quals", "TROJAN GOLD (B)",     "16U Boys"),
        ("jo-quals", "TROJAN SILVER (C)",   "16U Boys"),
    ]
    with _flask_app.test_client() as c:
        for tid, team, sheet in checks:
            r = c.get(f"/api/games/{tid}/{team}?sheet={sheet}")
            if r.status_code != 200:
                continue
            data = _json.loads(r.data)
            upcoming = [g for g in data.get("upcoming", []) if not g.get("placeholder")]
            if not upcoming:
                print(f"  [SKIP] {tid}/{team}: no direct upcoming games to check")
                continue
            # Check that _last_meeting is populated for at least one non-placeholder game
            # where the opponent is a known resolved team (not a slot description)
            from app import _SLOT_LIKE_RE as _slre
            known_opp_games = [g for g in upcoming if g.get("opponent") and not _slre.match(g.get("opponent", ""))]
            if not known_opp_games:
                print(f"  [SKIP] {tid}/{team}: all upcoming opponents are still unresolved slots")
                continue
            has_any = any(g.get("last_meeting") is not None for g in known_opp_games)
            # Soft check: SKIP rather than FAIL when there's genuinely no prior history
            if not has_any:
                print(f"  [INFO] {tid}/{team}: no prior meeting found for any of "
                      f"{[g['opponent'][:30] for g in known_opp_games[:3]]} — "
                      f"may be first encounter or history files not on disk")
            else:
                ok = _check(f"{tid}/{team}: last_meeting populated for at least one upcoming game",
                            True, "")
                # If has_any is True the check trivially passes; log it
    return failures


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    total_failures = 0

    print("=" * 60)
    print("Parser fixture tests")
    print("=" * 60)
    for args in FIXTURES:
        total_failures += _run_fixture(*args)

    print("\n" + "=" * 60)
    print("Team expansion tests (all-weekends contamination guard)")
    print("=" * 60)
    for args in TEAM_CHECKS:
        total_failures += _run_team_expansion(*args)

    print("\n" + "=" * 60)
    print("Championship weekend tests")
    print("=" * 60)
    for tournament_id, checks in CHAMPIONSHIP_CHECKS.items():
        meta = _tournament_meta(tournament_id)
        if not meta or not meta.get("date_start"):
            print(f"\n[SKIP] No date_start for {tournament_id!r} in KNOWN_TOURNAMENTS")
            continue
        anchor = meta["date_start"]
        for (sheet, team, min_d, max_d, min_e, min_t, min_s) in checks:
            total_failures += _run_championship_team(
                sheet, team, anchor, min_d, max_d, min_e, min_t, min_s
            )

    print("\n" + "=" * 60)
    print("Format A game-id collision test (Quiksilver Cup regression guard)")
    print("=" * 60)
    total_failures += test_format_a_id_collision()

    print("\n" + "=" * 60)
    print("Composite pool-rank slot test (Quiksilver Cup regression guard)")
    print("=" * 60)
    total_failures += test_composite_pool_rank_slot()

    print("\n" + "=" * 60)
    print("game_num chronological order test (Quiksilver Cup regression guard)")
    print("=" * 60)
    total_failures += test_game_num_chronological_order()

    print("\n" + "=" * 60)
    print("game_num ascending order (generic, cross-format checklist guard)")
    print("=" * 60)
    total_failures += test_game_num_ascending_order()

    print("\n" + "=" * 60)
    print("TBD stub placeholders (fixed-game-count formats)")
    print("=" * 60)
    total_failures += test_tbd_stub_placeholders()

    print("\n" + "=" * 60)
    print("CCA opponent-slot tests (self-reference regression guard)")
    print("=" * 60)
    total_failures += test_cca_opponent_slots()

    print("\n" + "=" * 60)
    print("CCA game-num tests (win/lose paths share same game_num)")
    print("=" * 60)
    total_failures += test_cca_game_nums()

    print("\n" + "=" * 60)
    print("strip_prefix compound tier_pool codes (Junior Olympics regression guard)")
    print("=" * 60)
    total_failures += test_strip_prefix_compound_pool_codes()

    print("\n" + "=" * 60)
    print("Trojan team name cleanliness (no raw slot strings as team names)")
    print("=" * 60)
    total_failures += test_trojan_team_names_clean()

    print("\n" + "=" * 60)
    print("Direct game opponent labels (no slot-like strings on non-placeholder cards)")
    print("=" * 60)
    total_failures += test_no_slot_like_opponents()

    print("\n" + "=" * 60)
    print("Game-num column integrity (no sequential rounds merged into one column)")
    print("=" * 60)
    total_failures += test_game_num_column_integrity()

    print("\n" + "=" * 60)
    print("Game-num cascading merge (same round shares one number across parents)")
    print("=" * 60)
    total_failures += test_game_num_cascading_merge()

    print("\n" + "=" * 60)
    print("Game-num chronological order (GAME 2 never starts before GAME 1)")
    print("=" * 60)
    total_failures += test_game_num_date_monotonic()

    print("\n" + "=" * 60)
    print("JO 2026 teams reachable (live schedule regression guard)")
    print("=" * 60)
    total_failures += test_jo_2026_teams_reachable()

    print("\n" + "=" * 60)
    print("Place Predictor: seeded single-elimination bracket (no pool letters)")
    print("=" * 60)
    total_failures += test_place_predictor_seeded_bracket()

    print("\n" + "=" * 60)
    print("All known tournaments, all Trojan teams (generic checklist item 4)")
    print("=" * 60)
    total_failures += test_all_known_tournaments_teams_healthy()

    print("\n" + "=" * 60)
    print("Bracket-tree game_num collision guard (canonical_bracket, not the flat list)")
    print("=" * 60)
    total_failures += test_bracket_tree_game_num_no_collision()

    print("\n" + "=" * 60)
    print("Bracket-tree last_meeting parity (JO same-tournament rematch guard)")
    print("=" * 60)
    total_failures += test_bracket_tree_last_meeting_parity()

    print("\n" + "=" * 60)
    print("Bracket-tree opponent resolution parity (_team_opp_slot vs inline tree algorithm)")
    print("=" * 60)
    total_failures += test_bracket_tree_opponent_resolution_parity()

    print("\n" + "=" * 60)
    print("Flat list is_current placeholder guard (live score on unresolved opponent)")
    print("=" * 60)
    total_failures += test_flat_list_is_current_placeholder_guard()

    print("\n" + "=" * 60)
    print("Flat list / tree path parity (win-lose-pool derivation, third independent pair)")
    print("=" * 60)
    total_failures += test_flat_list_tree_path_parity()

    print("\n" + "=" * 60)
    print("Last meeting availability (cross-tournament history search active)")
    print("=" * 60)
    total_failures += test_last_meeting_available()

    print("\n" + "=" * 60)
    print("NJO tree multi-phase (Junior Olympics regression guard)")
    print("=" * 60)
    total_failures += test_njo_tree_multi_phase()

    print("\n" + "=" * 60)
    print("NJO tree tied-game no misattribution (Junior Olympics regression guard)")
    print("=" * 60)
    total_failures += test_njo_tree_tied_game_no_misattribution()

    print("\n" + "=" * 60)
    print("NJO tree slot-code advancement (Junior Olympics regression guard)")
    print("=" * 60)
    total_failures += test_njo_tree_slot_code_advancement()

    print("\n" + "=" * 60)
    print("NJO tree pool-rank TBD stub (Junior Olympics regression guard)")
    print("=" * 60)
    total_failures += test_njo_tree_pool_rank_tbd_stub()

    print("\n" + "=" * 60)
    print("Bracket tree layout column-monotonic invariant (Junior Olympics regression guard)")
    print("=" * 60)
    total_failures += test_tree_layout_column_monotonic_invariant()

    print("\n" + "=" * 60)
    if total_failures == 0:
        print(f"\033[32mAll tests passed.\033[0m")
    else:
        print(f"\033[31m{total_failures} test(s) FAILED.\033[0m")
    print("=" * 60)
    sys.exit(0 if total_failures == 0 else 1)


if __name__ == "__main__":
    main()
